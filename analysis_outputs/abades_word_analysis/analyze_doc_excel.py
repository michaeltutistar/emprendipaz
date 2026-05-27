import json, zipfile, os, re
from pathlib import Path
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from openpyxl import load_workbook

base = Path(r"c:\Users\USUARIO\Documents\e-learning-platform\analysis_outputs\abades_word_analysis")
base.mkdir(parents=True, exist_ok=True)
docx_path = Path(r"C:\Users\USUARIO\Downloads\Informe_Metricas_nodo_Abades.docx")
xlsx_path = Path(r"C:\Users\USUARIO\Downloads\nuevos datos cordillera.xlsx")

# DOCX text + tables
summary = {"docx": str(docx_path), "exists": docx_path.exists(), "size": docx_path.stat().st_size if docx_path.exists() else None}
doc = Document(str(docx_path))
paragraphs = []
for i,p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt:
        style = p.style.name if p.style else None
        paragraphs.append({"index": i, "style": style, "text": txt})

tables = []
for ti,table in enumerate(doc.tables):
    rows=[]
    for row in table.rows:
        rows.append([cell.text.strip() for cell in row.cells])
    tables.append({"index": ti, "rows": rows})

# Extract images from docx zip
media_dir = base / "media"
media_dir.mkdir(exist_ok=True)
images=[]
with zipfile.ZipFile(docx_path, 'r') as z:
    for name in z.namelist():
        if name.startswith('word/media/'):
            out = media_dir / Path(name).name
            out.write_bytes(z.read(name))
            info={"file": str(out), "name": out.name, "bytes": out.stat().st_size}
            try:
                with Image.open(out) as im:
                    info.update({"width": im.width, "height": im.height, "mode": im.mode, "format": im.format})
            except Exception as e:
                info["error"] = str(e)
            images.append(info)

# Make image contact sheet
thumbs=[]
for info in images:
    try:
        im=Image.open(info['file']).convert('RGB')
        im.thumbnail((260,180))
        tile=Image.new('RGB',(300,230),'white')
        tile.paste(im,((300-im.width)//2,10))
        d=ImageDraw.Draw(tile)
        label=f"{info['name']}\n{info.get('width','?')}x{info.get('height','?')}"
        d.text((10,195),label,fill=(0,0,0))
        thumbs.append(tile)
    except Exception:
        pass
if thumbs:
    cols=2
    rows=(len(thumbs)+cols-1)//cols
    sheet=Image.new('RGB',(cols*300,rows*230),(240,240,240))
    for idx,t in enumerate(thumbs):
        sheet.paste(t,((idx%cols)*300,(idx//cols)*230))
    sheet.save(base/'image_contact_sheet.jpg', quality=90)

summary.update({"paragraph_count": len(paragraphs), "table_count": len(tables), "image_count": len(images), "images": images})
(base/'docx_summary.json').write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')

# Human readable docx text
with (base/'docx_text_outline.md').open('w', encoding='utf-8') as f:
    f.write(f"# Análisis texto DOCX\n\nArchivo: {docx_path}\n\n")
    f.write(f"Párrafos con texto: {len(paragraphs)}\nTablas: {len(tables)}\nImágenes: {len(images)}\n\n")
    f.write("## Párrafos\n")
    for p in paragraphs:
        f.write(f"- [{p['index']}] ({p['style']}) {p['text']}\n")
    f.write("\n## Tablas\n")
    for t in tables:
        f.write(f"\n### Tabla {t['index']} ({len(t['rows'])} filas)\n")
        for r in t['rows'][:20]:
            f.write(" | ".join(x.replace('\n',' / ') for x in r)+"\n")
        if len(t['rows'])>20: f.write(f"... {len(t['rows'])-20} filas más\n")

# Excel analysis
wb = load_workbook(xlsx_path, data_only=True)
excel = {"xlsx": str(xlsx_path), "sheets": []}
with (base/'excel_analysis.md').open('w', encoding='utf-8') as f:
    f.write(f"# Análisis Excel\n\nArchivo: {xlsx_path}\n\n")
    for ws in wb.worksheets:
        max_row, max_col = ws.max_row, ws.max_column
        rows=[]
        for row in ws.iter_rows(min_row=1, max_row=min(max_row, 20), values_only=True):
            rows.append(list(row))
        # detect used range non-empty
        non_empty=[]
        for row in ws.iter_rows(values_only=True):
            vals=list(row)
            if any(v is not None for v in vals):
                non_empty.append(vals)
        headers = non_empty[0] if non_empty else []
        sheet_info={"title": ws.title, "max_row": max_row, "max_col": max_col, "non_empty_rows": len(non_empty), "headers": headers, "sample_rows": non_empty[1:8]}
        excel["sheets"].append(sheet_info)
        f.write(f"## Hoja: {ws.title}\n")
        f.write(f"Rango: {max_row} filas x {max_col} cols; filas no vacías: {len(non_empty)}\n\n")
        if headers:
            f.write("Encabezados: "+" | ".join(str(h) for h in headers)+"\n\n")
        f.write("Primeras filas no vacías:\n")
        for r in non_empty[:12]:
            f.write(" | ".join('' if v is None else str(v) for v in r)+"\n")
        f.write("\n")
(base/'excel_summary.json').write_text(json.dumps(excel, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
print(base)
