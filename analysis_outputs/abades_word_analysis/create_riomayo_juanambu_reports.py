"""
Genera los informes de metricas para los nodos Rio Mayo y Juanambu siguiendo la
MISMA estructura y estilo del Informe_Metricas_nodo_Guambuyaco, pero leyendo los
datos frescos del export server-side (metrics_clean.json) en lugar del Excel.
"""
from __future__ import annotations

import json
import unicodedata
from pathlib import Path

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw

BASE = Path(__file__).resolve().parent
METRICS = Path(r"C:\Users\USUARIO\Documents\e-learning-platform\_lambda_patch\metrics_clean.json")
# Membrete real extraido del informe anterior (logos EmprendiPaz/DNP/SGR/etc.)
HEADER_SRC = Path(r"C:\Users\USUARIO\Documents\e-learning-platform\_informe_ref\word\media\image1.png")
DOWNLOADS = Path(r"C:\Users\USUARIO\Downloads")

GREEN = "16A34A"
PURPLE = "7E22CE"
DARK = "1F2937"

AWS_PLACEHOLDERS = [
    "CPUUtilization: Average",
    "DatabaseConnections: Sum",
    "FreeableMemory: Average",
    "ReadLatency: Average",
    "ReadThroughput: Average",
    "ReadIOPS: Average",
]

MODULE_SHORT = {
    "Atención al Cliente": "Atención al Cliente",
    "Descubrimiento de Oportunidades": "Descubr. Oportunidades",
    "Finanzas": "Finanzas",
    "Liderazgo": "Liderazgo",
    "Marketing Digital": "Mkt. Digital",
    "Marketing y Comercialización": "Mkt. y Comercializ.",
    "Modelo de Negocios": "Modelo de Negocios",
    "Plan de Inversión": "Plan de Inversión",
    "Proyecto de vida": "Proyecto de Vida",
    "Trabajo en Equipo": "Trabajo en Equipo",
}

NODES = [
    {
        "slug": "rio_mayo",
        "name": "Río Mayo",
        "central": "La Cruz",
        "municipios": ["Albán", "Belén", "Colón", "El Tablón de Gómez", "La Cruz", "San Bernardo", "San Pablo"],
        "out": DOWNLOADS / "Informe_Metricas_nodo_Rio_Mayo.docx",
    },
    {
        "slug": "juanambu",
        "name": "Juanambú",
        "central": "La Unión",
        "municipios": ["Arboleda", "Buesaco", "La Unión", "San Lorenzo", "San Pedro de Cartago"],
        "out": DOWNLOADS / "Informe_Metricas_nodo_Juanambu.docx",
    },
]


def norm(value: str) -> str:
    s = (value or "").strip().lower()
    s = unicodedata.normalize("NFD", s)
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def fmt_num(value: float) -> str:
    return f"{value:.2f}".replace(".", ",")


def fmt_pct_num(value: float) -> str:
    return f"{float(value):.2f}%".replace(".", ",")


def performance(avg: float) -> str:
    if 30 <= avg < 40:
        return "Alto"
    if 40 <= avg < 50:
        return "Medio"
    if avg >= 50:
        return "Bajo"
    return "Sin clasificar"


def load_metrics() -> dict:
    data = json.loads(METRICS.read_text(encoding="utf-8"))
    resumen = {norm(r["municipio"]): r for r in data["resumen_por_municipio"]}
    intentos = {norm(r["municipio"]): r for r in data["intentos_por_modulo"]}
    modules = data["modules_order"]
    return {"resumen": resumen, "intentos": intentos, "modules": modules}


def build_node_data(municipios: list[str], metrics: dict) -> dict:
    resumen_map = metrics["resumen"]
    intentos_map = metrics["intentos"]
    modules = metrics["modules"]

    municipality_rows = []
    module_by_muni = []
    module_totals = {m: 0 for m in modules}

    for muni in municipios:
        key = norm(muni)
        r = resumen_map.get(key)
        if not r:
            continue
        estudiantes = int(r["total_estudiantes"])
        intentos = int(r["intentos_totales"])
        promedio = round(intentos / estudiantes, 2) if estudiantes else 0.0
        municipality_rows.append({
            "Municipio": muni,
            "Estudiantes": estudiantes,
            "Intentos": intentos,
            "Promedio": promedio,
            "Rendimiento": performance(promedio),
            "Finalización": fmt_pct_num(r["pct_finalizacion_modulos"]),
            "Asistencia": fmt_pct_num(r["pct_asistencia"]),
        })

        im = intentos_map.get(key, {})
        item = {"Municipio": muni}
        for m in modules:
            v = int(im.get(m, 0) or 0)
            item[m] = v
            module_totals[m] += v
        module_by_muni.append(item)

    total_students = sum(r["Estudiantes"] for r in municipality_rows)
    total_attempts = sum(r["Intentos"] for r in municipality_rows)
    avg_attempts = round(total_attempts / total_students, 2) if total_students else 0.0

    fin_values = [float(str(r["Finalización"]).replace("%", "").replace(",", ".")) for r in municipality_rows]
    asis_values = [float(str(r["Asistencia"]).replace("%", "").replace(",", ".")) for r in municipality_rows]
    avg_fin = round(sum(fin_values) / len(fin_values), 2) if fin_values else 0.0
    avg_asis = round(sum(asis_values) / len(asis_values), 2) if asis_values else 0.0

    ranking = sorted(module_totals.items(), key=lambda x: x[1], reverse=True)

    return {
        "municipality_rows": municipality_rows,
        "total_students": total_students,
        "total_attempts": total_attempts,
        "avg_attempts": avg_attempts,
        "general_performance": performance(avg_attempts),
        "avg_completion": f"{avg_fin:.2f}%".replace(".", ","),
        "avg_attendance": f"{avg_asis:.2f}%".replace(".", ","),
        "modules": modules,
        "module_by_muni": module_by_muni,
        "module_totals": module_totals,
        "ranking": ranking,
    }


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(9)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor.from_string(DARK if level == 1 else PURPLE)
    return paragraph


def add_table(doc: Document, headers, data_rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, header in enumerate(headers):
        set_cell_text(hdr[i], header, bold=True, color="FFFFFF")
        set_cell_shading(hdr[i], PURPLE)
    for row in data_rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = width
    doc.add_paragraph()
    return table


def make_municipality_chart(node_name: str, rows: list[dict], out: Path) -> Path:
    names = [r["Municipio"] for r in rows]
    totals = [r["Intentos"] for r in rows]
    colors = ["#14532d", "#16a34a", "#4ade80", "#6b21a8", "#7e22ce", "#c4b5fd", "#a855f7"]
    # repetir paleta si hay mas municipios
    while len(colors) < len(names):
        colors += colors
    fig, ax = plt.subplots(figsize=(10, 5))
    bars = ax.bar(names, totals, color=colors[: len(names)], edgecolor="white", linewidth=0.8)
    ax.set_title(f"Nodo {node_name} - Intentos totales por municipio", fontsize=14, fontweight="bold")
    ax.set_ylabel("Intentos totales")
    ymax = max(totals) if totals and max(totals) > 0 else 10
    ax.set_ylim(0, ymax * 1.25 + 1)
    ax.grid(axis="y", alpha=0.25)
    plt.xticks(rotation=20, ha="right")
    for bar, row in zip(bars, rows):
        height = bar.get_height()
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            height + (ymax * 0.02 if ymax else 0.2),
            f"{row['Intentos']}\nProm. {fmt_num(row['Promedio'])}",
            ha="center",
            va="bottom",
            fontsize=9,
        )
    fig.tight_layout()
    fig.savefig(out, dpi=150, bbox_inches="tight")
    plt.close(fig)
    return out


def make_modules_chart(node_name: str, ranking: list[tuple[str, int]], out: Path) -> Path:
    ordered = list(reversed(ranking))
    modules = [m for m, _ in ordered]
    totals = [t for _, t in ordered]
    colors = ["#065f46"] * 3 + ["#ea580c"] + ["#16a34a"] * 4 + ["#14532d"] * 3
    colors = colors[-len(modules):]

    fig, ax = plt.subplots(figsize=(10, 6))
    bars = ax.barh(modules, totals, color=colors, edgecolor="white", linewidth=0.8)
    ax.set_title(f"Nodo {node_name} - Ranking de intentos por módulo", fontsize=14, fontweight="bold")
    ax.set_xlabel("Intentos totales")
    xmax = max(totals) if totals and max(totals) > 0 else 10
    ax.set_xlim(0, xmax * 1.2 + 1)
    ax.grid(axis="x", alpha=0.25)
    for bar, total in zip(bars, totals):
        ax.text(bar.get_width() + 0.3, bar.get_y() + bar.get_height() / 2, str(total), va="center", fontsize=9)
    fig.tight_layout()
    fig.savefig(out, dpi=150, bbox_inches="tight")
    plt.close(fig)
    return out


def make_aws_placeholder(label: str, out: Path) -> Path:
    width, height = 1780, 730
    image = Image.new("RGB", (width, height), "#f3f4f6")
    draw = ImageDraw.Draw(image)
    draw.rectangle((30, 30, width - 30, height - 30), outline="#9ca3af", width=4)
    text = (
        "ESPACIO RESERVADO PARA CAPTURA AWS CLOUDWATCH\n"
        f"{label}\n"
        "(Pegar aquí el recorte correspondiente al periodo del informe)"
    )
    draw.multiline_text((80, height // 2 - 60), text, fill="#374151", spacing=10)
    image.save(out)
    return out


def municipios_list_text(municipios: list[str]) -> str:
    if len(municipios) == 1:
        return municipios[0]
    return ", ".join(municipios[:-1]) + " y " + municipios[-1]


def build_observations(node_name, rows, total_students, total_attempts):
    if not rows:
        return ["No se registraron municipios para este nodo en el export de estadísticas."]
    max_students = max(rows, key=lambda r: r["Estudiantes"])
    max_attempts = max(rows, key=lambda r: r["Intentos"])
    max_prom = max(rows, key=lambda r: r["Promedio"])
    min_prom = min(rows, key=lambda r: r["Promedio"])
    return [
        (
            f"{max_students['Municipio']} concentra el mayor número de estudiantes del nodo "
            f"({max_students['Estudiantes']} de {total_students}) y registra {max_students['Intentos']} intentos, "
            f"con un promedio de {fmt_num(max_students['Promedio'])} intentos por estudiante y rendimiento {max_students['Rendimiento']}."
        ),
        (
            f"{max_attempts['Municipio']} registra el mayor total de intentos del nodo ({max_attempts['Intentos']}), "
            f"con un promedio de {fmt_num(max_attempts['Promedio'])} intentos por estudiante."
        ),
        (
            f"{min_prom['Municipio']} evidencia el promedio de intentos más bajo del nodo ({fmt_num(min_prom['Promedio'])}), "
            f"con finalización de {min_prom['Finalización']} y asistencia de {min_prom['Asistencia']}, "
            "lo que sugiere una ruta de aprendizaje relativamente más eficiente dentro del grupo."
        ),
    ]


def build_document(node_cfg: dict, data: dict) -> Path:
    node_name = node_cfg["name"]
    municipios = node_cfg["municipios"]
    slug = node_cfg["slug"]
    out_docx = node_cfg["out"]

    rows = data["municipality_rows"]
    total_students = data["total_students"]
    total_attempts = data["total_attempts"]
    avg_attempts = data["avg_attempts"]
    general_perf = data["general_performance"]
    modules = data["modules"]
    module_by_muni = data["module_by_muni"]
    module_totals = data["module_totals"]
    ranking = data["ranking"]

    chart_muni = make_municipality_chart(node_name, rows, BASE / f"{slug}_intentos_municipio.png")
    chart_modules = make_modules_chart(node_name, ranking, BASE / f"{slug}_ranking_modulos.png")

    hardest = ranking[:3] if ranking else []
    easiest = list(reversed(ranking[-3:])) if ranking else []
    perf_label = (
        f"{general_perf} (30–39,9)" if general_perf == "Alto"
        else (f"{general_perf} (40–49,9)" if general_perf == "Medio"
              else (f"{general_perf} (>50)" if general_perf == "Bajo" else general_perf))
    )

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)

    if HEADER_SRC.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(HEADER_SRC), width=Inches(6.7))

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("INFORME DE ESTADÍSTICAS GENERALES")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor.from_string(DARK)

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"NODO {node_name.upper()} — PLATAFORMA EMPRENDIPAZ")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string(PURPLE)

    doc.add_paragraph()
    add_heading(doc, "1. Resumen General del Nodo", 1)
    doc.add_paragraph(
        f"A continuación, se presentan los indicadores consolidados para los {total_students} estudiantes del Nodo {node_name}, "
        f"con base en las métricas de intentos, finalización de módulos y asistencia registradas en la plataforma Emprendipaz. "
        f"Este nodo integra {len(municipios)} municipios del departamento de Nariño: {municipios_list_text(municipios)}."
    )

    add_table(
        doc,
        ["INDICADOR", "VALOR"],
        [
            ["Total de estudiantes en el nodo", total_students],
            ["Total de municipios participantes", len(rows)],
            ["Total de intentos registrados", total_attempts],
            ["Promedio de intentos por estudiante", fmt_num(avg_attempts)],
            ["Rendimiento general según intentos", perf_label],
            ["% Finalización de módulos (promedio nodo)", data["avg_completion"]],
            ["% Asistencia general del nodo", data["avg_attendance"]],
        ],
        widths=[Inches(3.5), Inches(2.2)],
    )

    paragraph = doc.add_paragraph()
    paragraph.add_run("Interpretación:").bold = True
    doc.add_paragraph(
        f"El Nodo {node_name} alcanzó un promedio de {fmt_num(avg_attempts)} intentos por estudiante, "
        f"ubicándose en rendimiento {general_perf}. La finalización promedio del nodo es {data['avg_completion']} "
        f"y la asistencia promedio es {data['avg_attendance']}. En total, los {total_students} estudiantes generaron "
        f"{total_attempts} intentos en la plataforma."
    )

    add_heading(doc, "2. Análisis por Municipio", 1)
    doc.add_paragraph(
        "La siguiente tabla presenta el desglose del desempeño por municipio, incluyendo número de estudiantes, intentos totales y promedio, "
        "rendimiento según la escala de la plataforma, porcentaje de finalización de módulos y porcentaje de asistencia."
    )
    doc.add_paragraph("*Rendimiento según intentos: de 30–39,9: Alto | de 40–49,9: Medio | >50: Bajo | <30: Sin clasificar")

    muni_table = [
        [r["Municipio"], r["Estudiantes"], r["Intentos"], fmt_num(r["Promedio"]), r["Rendimiento"], r["Finalización"], r["Asistencia"]]
        for r in rows
    ]
    muni_table.append(["TOTALES", total_students, total_attempts, fmt_num(avg_attempts), general_perf, data["avg_completion"], data["avg_attendance"]])
    add_table(
        doc,
        ["Municipio", "N° Est.", "Intentos Totales", "Promedio Intentos / Est.", "Rendimiento", "% Finalización", "% Asistencia"],
        muni_table,
    )

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(chart_muni), width=Inches(6.4))

    add_heading(doc, "Observaciones por municipio", 2)
    for text in build_observations(node_name, rows, total_students, total_attempts):
        doc.add_paragraph(text)

    add_heading(doc, "3. Conclusiones", 1)
    doc.add_paragraph(
        f"El Nodo {node_name} demuestra un desempeño general sólido dentro de la plataforma Emprendipaz, "
        f"con un promedio de {fmt_num(avg_attempts)} intentos por estudiante (rendimiento {general_perf}) "
        f"y finalización promedio de {data['avg_completion']}."
    )
    doc.add_paragraph(
        "La distribución de intentos muestra diferencias territoriales moderadas entre municipios, "
        "con participación activa en el proceso formativo."
    )
    if hardest:
        doc.add_paragraph(
            "Se recomienda mantener estrategias de acompañamiento pedagógico focalizadas en los módulos con mayor número de intentos, "
            f"especialmente {hardest[0][0]}"
            + (f", {hardest[1][0]} y {hardest[2][0]}." if len(hardest) >= 3 else ".")
        )

    add_heading(doc, "4. Análisis de Intentos por Módulo", 1)
    doc.add_paragraph(
        f"La siguiente tabla presenta el total de intentos registrados por módulo en cada municipio del Nodo {node_name}. "
        "Este análisis permite identificar cuáles contenidos generan mayor dificultad y cuáles son apropiados con mayor facilidad."
    )

    headers = ["Municipio"] + [MODULE_SHORT.get(m, m) for m in modules]
    module_rows = [[row["Municipio"]] + [row[m] for m in modules] for row in module_by_muni]
    module_rows.append(["TOTALES"] + [module_totals[m] for m in modules])
    add_table(doc, headers, module_rows)

    if hardest:
        top_name, top_total = hardest[0]
        add_heading(doc, f"4.1. Módulo con Mayor Dificultad: {top_name}", 2)
        doc.add_paragraph(
            f"El módulo de {top_name} registró el mayor número de intentos totales en el nodo, con {top_total} intentos, "
            "posicionándose como el contenido de mayor dificultad para los participantes."
        )
        easy_name, easy_total = ranking[-1]
        add_heading(doc, f"4.2. Módulo con Mayor Facilidad: {easy_name}", 2)
        doc.add_paragraph(
            f"El módulo de {easy_name} registró el menor total de intentos del nodo, con {easy_total} intentos, "
            "lo que indica una apropiación más fluida del contenido."
        )

    add_heading(doc, "4.3. Ranking de Módulos por Intentos", 2)
    ranking_rows = []
    for idx, (module, total) in enumerate(ranking, start=1):
        if idx == 1:
            nivel = "Mayor dificultad"
        elif idx <= 3:
            nivel = "Alta dificultad"
        elif idx >= len(ranking) - 2:
            nivel = "Menor dificultad"
        else:
            nivel = "Dificultad media"
        ranking_rows.append([idx, module, total, nivel])
    add_table(
        doc,
        ["Pos.", "Módulo", "Total Intentos", "Nivel"],
        ranking_rows,
        widths=[Inches(0.6), Inches(3.3), Inches(1.2), Inches(1.6)],
    )

    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(chart_modules), width=Inches(6.4))

    add_heading(doc, "5. Métricas Generales del nodo", 1)
    doc.add_paragraph(f"Plataforma Emprendipaz — Nodo {node_name}")
    add_table(
        doc,
        ["Métrica", "Resultado"],
        [
            ["Municipios participantes", municipios_list_text(municipios)],
            ["Total de estudiantes", total_students],
            ["Intentos totales", total_attempts],
            ["Promedio general de intentos", fmt_num(avg_attempts)],
            ["Rendimiento general", general_perf],
            ["Finalización de módulos (promedio)", data["avg_completion"]],
            ["Asistencia general (promedio)", data["avg_attendance"]],
        ],
    )

    for idx, label in enumerate(AWS_PLACEHOLDERS, start=1):
        placeholder = make_aws_placeholder(label, BASE / f"{slug}_aws_{idx}.png")
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(placeholder), width=Inches(6.5))

    doc.add_paragraph(f"Plataforma Emprendipaz — Nodo {node_name}")

    add_heading(doc, "CONCLUSIONES FINALES", 1)
    top3 = ", ".join(m for m, _ in hardest[:3]) if hardest else "los módulos identificados"
    easy3 = ", ".join(m for m, _ in easiest) if easiest else "los módulos de menor esfuerzo"
    final_sections = [
        ("1. Desempeño general del nodo",
         f"El Nodo {node_name} evidencia un promedio de {fmt_num(avg_attempts)} intentos por estudiante "
         f"(rendimiento {general_perf}), con {total_attempts} intentos totales, finalización promedio de {data['avg_completion']} "
         f"y asistencia promedio de {data['avg_attendance']}."),
        ("2. Compromiso de los participantes",
         "Los resultados reflejan participación activa en la plataforma y avance en el proceso formativo en la mayoría de los municipios del nodo."),
        ("3. Diferenciación territorial",
         build_observations(node_name, rows, total_students, total_attempts)[0]),
        ("4. Módulos críticos de aprendizaje",
         f"{top3} concentran los mayores intentos; se recomienda reforzar acompañamiento en esos contenidos."),
        ("5. Módulos de mayor facilidad",
         f"{easy3} presentan menor número de intentos, lo que sugiere una apropiación más fluida."),
        ("6. Infraestructura AWS",
         "Pegar en la sección 5 las capturas de CloudWatch (CPU, conexiones, memoria, latencia, throughput e IOPS) en los espacios reservados."),
        ("7. Conclusión estratégica",
         f"El Nodo {node_name} muestra condiciones favorables para consolidar competencias emprendedoras "
         "y replicar buenas prácticas dentro del programa Emprendipaz."),
    ]
    for title, body in final_sections:
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(PURPLE)
        doc.add_paragraph(body)

    doc.save(out_docx)
    return out_docx


def main():
    metrics = load_metrics()
    for node in NODES:
        data = build_node_data(node["municipios"], metrics)
        out = build_document(node, data)
        summary = BASE / f"{node['slug']}_computed_summary.json"
        summary.write_text(json.dumps({"node": node["name"], "output": str(out), **data}, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"OK {node['name']}: {data['total_students']} est, {data['total_attempts']} intentos -> {out}")


if __name__ == "__main__":
    main()
