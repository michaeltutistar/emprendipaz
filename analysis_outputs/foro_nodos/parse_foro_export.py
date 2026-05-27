"""Parsea export Excel del foro por nodo (hojas Resumen y Conversaciones)."""
from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

MUNICIPIO_ALIASES = {
    "El Penol": "El Peñol",
    "Imues": "Imués",
    "Tuquerres": "Túquerres",
}


def norm_municipio(value: str | None) -> str:
    raw = str(value or "").strip()
    if not raw:
        return "Sin municipio"
    return MUNICIPIO_ALIASES.get(raw, raw)


@dataclass
class ForumReply:
    author: str
    municipio: str
    body: str
    created_hint: str = ""


@dataclass
class ForumThread:
    thread_id: int
    question: str
    status: str
    replies: list[ForumReply] = field(default_factory=list)

    @property
    def student_replies(self) -> list[ForumReply]:
        return self.replies


@dataclass
class ForumNodeData:
    name: str
    slug: str
    central: str
    municipios: list[str]
    total_threads: int
    total_replies: int
    exported_at: str
    threads: list[ForumThread] = field(default_factory=list)

    @property
    def main_thread(self) -> ForumThread | None:
        if not self.threads:
            return None
        return max(self.threads, key=lambda t: len(t.replies))


def parse_foro_xlsx(path: Path, slug: str) -> ForumNodeData:
    wb = load_workbook(path, read_only=True, data_only=True)
    resumen = {
        str(row[0]).strip(): row[1]
        for row in wb["Resumen"].iter_rows(values_only=True)
        if row and row[0] and str(row[0]) != "Campo"
    }
    rows = list(wb["Conversaciones"].iter_rows(values_only=True))
    wb.close()

    I_HILO, I_STATUS, I_Q, I_TIPO, I_ROL, I_MAUT, I_MSG, I_FECHA = (
        2,
        3,
        8,
        9,
        11,
        12,
        15,
        13,
    )

    hilos_raw: dict[int, list] = defaultdict(list)
    for row in rows[1:]:
        if row[I_HILO] is None:
            continue
        hilos_raw[int(row[I_HILO])].append(row)

    threads: list[ForumThread] = []
    for hid, msgs in sorted(hilos_raw.items(), key=lambda x: -len(x[1])):
        question = str(msgs[0][I_Q] or "").strip()
        status = str(msgs[0][I_STATUS] or "").strip()
        thread = ForumThread(thread_id=hid, question=question, status=status)
        for row in msgs:
            if str(row[I_TIPO]) != "Respuesta":
                continue
            if str(row[I_ROL] or "").lower() != "usuario":
                continue
            municipio = norm_municipio(row[I_MAUT])
            if municipio in ("Pasto", "Samaniego"):
                continue
            body = str(row[I_MSG] or "").strip()
            if not body:
                continue
            thread.replies.append(
                ForumReply(
                    author=str(row[10] or "Estudiante").strip(),
                    municipio=municipio,
                    body=body,
                    created_hint=str(row[I_FECHA] or ""),
                )
            )
        if thread.replies or question:
            threads.append(thread)

    municipios = [
        m.strip()
        for m in str(resumen.get("Municipios", "")).split(",")
        if m.strip()
    ]

    return ForumNodeData(
        name=str(resumen.get("Nodo", slug)).strip(),
        slug=slug,
        central=str(resumen.get("Cabecera", "")).strip(),
        municipios=municipios,
        total_threads=int(resumen.get("Total hilos", len(threads)) or 0),
        total_replies=int(resumen.get("Total respuestas", 0) or 0),
        exported_at=f"{resumen.get('Exportado el', '')} {resumen.get('Exportado a las', '')}".strip(),
        threads=threads,
    )


def municipio_counts(data: ForumNodeData, thread: ForumThread | None = None) -> list[tuple[str, int]]:
    counter: Counter[str] = Counter()
    threads = [thread] if thread else data.threads
    for t in threads:
        for reply in t.replies:
            counter[reply.municipio] += 1
    return sorted(counter.items(), key=lambda x: (-x[1], x[0]))


def pick_quotes(thread: ForumThread, limit: int = 4) -> list[ForumReply]:
    seen_munis: set[str] = set()
    picks: list[ForumReply] = []
    for reply in sorted(thread.replies, key=lambda r: -len(r.body)):
        if reply.municipio in seen_munis and len(picks) >= 2:
            continue
        if len(reply.body) < 80:
            continue
        picks.append(reply)
        seen_munis.add(reply.municipio)
        if len(picks) >= limit:
            break
    return picks
