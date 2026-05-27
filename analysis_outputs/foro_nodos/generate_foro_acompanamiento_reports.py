"""
Genera informes de acompañamiento del foro (solo foro, sin métricas)
para Guambuyaco y Sabana, en Word separados.
"""
from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

from parse_foro_export import (
    ForumNodeData,
    municipio_counts,
    parse_foro_xlsx,
    pick_quotes,
)

DOWNLOADS = Path(r"C:\Users\USUARIO\Downloads")
TEMPLATE = DOWNLOADS / "Informe Acompañamiento Foro Exprovincia de Obando.docx"
CAPTURES = Path(r"C:\Users\USUARIO\Downloads\Abril-Emprendipaz\capturas_foro")
PERIODO = "01/05/2026 a 26/05/2026"

NODES = [
    {
        "slug": "guambuyaco",
        "xlsx": DOWNLOADS / "foro-guambuyaco-2026-05-26.xlsx",
        "output": DOWNLOADS / "Informe_acompañamiento_foro_Guambuyaco.docx",
    },
    {
        "slug": "sabana",
        "xlsx": DOWNLOADS / "foro-sabana-2026-05-26.xlsx",
        "output": DOWNLOADS / "Informe_acompañamiento_foro_Sabana.docx",
    },
]


def clear_document_body(document: Document) -> None:
    body = document._element.body
    section_properties = body.sectPr
    for child in list(body):
        if child is not section_properties:
            body.remove(child)


def add_center(document: Document, text: str, style: str = "Title") -> None:
    p = document.add_paragraph()
    try:
        p.style = style
    except KeyError:
        p.style = "Normal"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    style = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}.get(level, "Heading 2")
    p = document.add_paragraph()
    try:
        p.style = style
    except KeyError:
        p.style = "Normal"
    p.add_run(text)


def add_para(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.style = "Normal"
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.add_run(text)


def add_bullet(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.style = "Normal"
    p.add_run(f"• {text}")


def add_table(document: Document, headers: list[str], rows: list[list]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        for run in table.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    document.add_paragraph()


def add_image(document: Document, path: Path, caption: str, width_in: float = 6.2) -> None:
    if not path.exists():
        add_para(document, f"[Pendiente captura: {caption}]")
        return
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width_in))
    cap = document.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(caption)


def build_report(data: ForumNodeData, captures_dir: Path, output_path: Path) -> Document:
    if TEMPLATE.exists():
        shutil.copy2(TEMPLATE, output_path)
        doc = Document(str(output_path))
        clear_document_body(doc)
    else:
        doc = Document()

    main = data.main_thread
    main_replies = len(main.replies) if main else 0
    active_threads = [t for t in data.threads if t.replies]
    muni_all = municipio_counts(data)
    muni_main = municipio_counts(data, main) if main else []

    add_center(doc, "IMPLEMENTACIÓN DEL FORO")
    add_center(doc, f"Periodo: {PERIODO}", "Subtitle")
    doc.add_paragraph()
    add_heading(doc, f"FORO EN EL NODO {data.name.upper()}")
    add_para(
        doc,
        f"De acuerdo con el reporte exportado desde la plataforma EmprendiPaz ({data.exported_at}) "
        f"y la evidencia revisada en el dashboard del tutor (Foros por nodo), el nodo {data.name} "
        f"(cabecera {data.central}) registró {data.total_threads} hilos y {data.total_replies} "
        f"intervenciones de respuesta en el periodo. El análisis siguiente corresponde exclusivamente "
        f"al componente de foro y acompañamiento territorial, sin incluir métricas de intentos ni "
        f"desempeño por módulo.",
    )

    add_heading(doc, "ANÁLISIS DE INTERACCIÓN DEL NODO", 2)
    add_para(doc, "Este comportamiento se refleja en los siguientes indicadores:")
    add_bullet(doc, f"Número de hilos abiertos: {data.total_threads}")
    add_bullet(doc, f"Número total de respuestas (exportación): {data.total_replies}")
    if main:
        add_bullet(doc, f"Respuestas de estudiantes en el hilo principal: {main_replies}")
        add_bullet(doc, f"Municipios del nodo con participación en el hilo principal: {len(muni_main)}")
    add_bullet(doc, f"Hilos con participación efectiva: {len(active_threads)}")
    nivel = "Alto" if main_replies >= 30 else "Medio" if main_replies >= 15 else "Inicial"
    add_bullet(doc, f"Nivel de interacción: {nivel}")

    add_heading(doc, "ANÁLISIS TÉCNICO Y FUNCIONAL", 2)
    add_heading(doc, "1. Caracterización general del foro", 3)
    add_para(
        doc,
        f"Se identifica un hilo principal (ID {main.thread_id if main else 'N/A'}) con "
        f"{main_replies} respuestas de estudiantes del nodo, lo que evidencia participación "
        f"activa orientada a la aplicación práctica del contenido formativo. "
        f"Adicionalmente se registran {max(0, len(active_threads) - 1)} hilos complementarios "
        f"sobre herramientas de publicidad, alianzas, fidelización y reflexión sobre el proceso formativo.",
    )

    add_heading(doc, "2. Temáticas predominantes", 3)
    themes = [
        "Acciones concretas en 7 días para potenciar o mejorar el emprendimiento (consigna principal del tutor).",
        "Marketing digital, redes sociales, publicaciones creativas y mayor visibilidad comercial.",
        "Estrategias de bajo costo para dar a conocer el negocio y captar clientes.",
        "Fidelización de clientes y cierre de ventas en contacto directo.",
        "Reflexión sobre utilidad del proceso formativo y proyección del emprendimiento.",
    ]
    for i, theme in enumerate(themes, 1):
        add_bullet(doc, f"{i}. {theme}")

    add_heading(doc, "3. Participación por municipio", 3)
    add_para(
        doc,
        "La siguiente tabla resume la distribución de respuestas de estudiantes (rol usuario) "
        f"en el nodo {data.name}, excluyendo intervenciones del equipo tutor:"
    )
    rows = [[m, c] for m, c in muni_all]
    add_table(doc, ["Municipio", "Número de respuestas"], rows)
    if muni_all:
        top = muni_all[0]
        add_para(
            doc,
            f"La mayor concentración de respuestas se observa en {top[0]} ({top[1]} intervenciones). "
            f"Participaron {len(muni_all)} de los {len(data.municipios)} municipios del nodo en el periodo analizado.",
        )

    if main:
        add_heading(doc, "4. Pregunta orientadora del tutor (hilo principal)", 3)
        add_para(doc, "La consigna principal publicada en el foro fue:")
        q = document_quote(main.question)
        add_para(doc, f"«{q}»")
        add_heading(doc, "5. Respuestas representativas", 3)
        add_para(
            doc,
            "A continuación se citan intervenciones textuales de estudiantes, seleccionadas por claridad "
            "y diversidad municipal:",
        )
        for reply in pick_quotes(main, 4):
            body = document_quote(reply.body[:500])
            add_para(doc, f"«{body}»")
            add_para(doc, f"— {reply.author} ({reply.municipio})")

    add_heading(doc, "6. Evidencia gráfica (capturas de la plataforma)", 3)
    add_image(
        doc,
        captures_dir / "01_lista_hilos_nodo.png",
        f"Imagen 1. Vista del foro — nodo {data.name} (listado de hilos)",
    )
    add_image(
        doc,
        captures_dir / "02_hilo_principal_respuestas.png",
        f"Imagen 2. Hilo principal con respuestas de estudiantes — nodo {data.name}",
    )

    add_heading(doc, "7. Conclusiones", 2)
    add_para(
        doc,
        f"El foro del nodo {data.name} operó como un espacio efectivo de acompañamiento y aplicación práctica. "
        f"La dinámica del tutor movilizó respuestas concretas vinculadas a visibilidad, marketing, servicio al cliente "
        f"y acciones de corto plazo en los emprendimientos. La participación multisectorial por municipios confirma "
        f"apropiación territorial del espacio formativo.",
    )
    add_para(
        doc,
        "Se recomienda mantener la consigna de acciones en 7 días, reforzar seguimiento en municipios con menor "
        f"número de intervenciones y continuar usando el exportador del dashboard del tutor como evidencia "
        f"auditable del acompañamiento.",
    )
    return doc


def document_quote(text: str) -> str:
    return " ".join((text or "").replace("\u200b", " ").split())


def main() -> None:
    for cfg in NODES:
        data = parse_foro_xlsx(cfg["xlsx"], cfg["slug"])
        captures = CAPTURES / cfg["slug"]
        doc = build_report(data, captures, cfg["output"])
        doc.save(str(cfg["output"]))
        print("Generado:", cfg["output"])


if __name__ == "__main__":
    main()
