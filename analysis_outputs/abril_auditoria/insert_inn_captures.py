"""Inserta capturas INN en INN ABR_CORREGIDO.docx (reemplaza figuras 1-3)."""
from pathlib import Path

from docx import Document
from docx.shared import Inches

DOC = Path(r"C:\Users\USUARIO\Downloads\Abril-Emprendipaz\INN ABR_CORREGIDO.docx")
CAP = Path(r"C:\Users\USUARIO\Downloads\Abril-Emprendipaz\capturas_inn")

# Párrafos con dibujo embebido (índice detectado en el docx corregido)
IMAGE_PARAS = {
    1: (154, CAP / "01_modo_online_dashboard.png"),
    2: (162, CAP / "03_modo_offline_network.png"),
    3: (177, CAP / "04_reconexion_online.png"),
}


def replace_paragraph_image(paragraph, image_path: Path, width_in: float = 6.2) -> None:
    p = paragraph._element
    for child in list(p):
        tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
        if tag in ("r", "drawing", "pict"):
            p.remove(child)
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_in))


def main() -> None:
    doc = Document(str(DOC))
    for num, (idx, img) in IMAGE_PARAS.items():
        if not img.exists():
            raise FileNotFoundError(f"Falta captura Imagen {num}: {img}")
        replace_paragraph_image(doc.paragraphs[idx], img)
    doc.save(str(DOC))
    print("Actualizado:", DOC)
    for n, (_, p) in IMAGE_PARAS.items():
        print(f"  Imagen {n} <- {p.name}")


if __name__ == "__main__":
    main()
