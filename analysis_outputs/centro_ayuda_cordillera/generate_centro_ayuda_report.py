import json
import subprocess
import collections
import statistics
import csv
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import matplotlib.pyplot as plt

ROOT = Path(r"c:\Users\USUARIO\Documents\e-learning-platform")
OUT = ROOT / "analysis_outputs" / "centro_ayuda_cordillera"
OUT.mkdir(parents=True, exist_ok=True)
DOWNLOADS = Path(r"C:\Users\USUARIO\Downloads")
BUCKET = "elearning-archivos"
LATEST_PREFIX = "support/conversations/latest/"
HISTORY_PREFIX = "support/conversations/history/"
DOCX_OUT = DOWNLOADS / "Informe_Centro_Ayuda_Nodo_Cordillera.docx"
XLSX_OUT = DOWNLOADS / "Consolidado_Centro_Ayuda_Cordillera.xlsx"
SUMMARY_JSON_OUT = OUT / "centro_ayuda_cordillera_summary.json"
MASTER_CSV_OUT = OUT / "centro_ayuda_cordillera_tickets.csv"
MESSAGES_CSV_OUT = OUT / "centro_ayuda_cordillera_messages.csv"
WHATSAPP_EVIDENCE_DIR = OUT / "evidence" / "whatsapp"
WHATSAPP_EVIDENCE = [
    {
        "ticket_ids": [31],
        "path": WHATSAPP_EVIDENCE_DIR / "ticket_31.png",
        "description": "Mensaje inicial enviado a WhatsApp con ticket #31, usuario y resumen generado desde la plataforma.",
    },
    {
        "ticket_ids": [30],
        "path": WHATSAPP_EVIDENCE_DIR / "ticket_30.png",
        "description": "Mensaje inicial enviado a WhatsApp con ticket #30, usuario y resumen del inconveniente en foro.",
    },
    {
        "ticket_ids": [10, 5],
        "path": WHATSAPP_EVIDENCE_DIR / "ticket_10_5.png",
        "description": "Captura de conversación con evidencia del ticket #10 y mensaje inicial del ticket #5.",
    },
    {
        "ticket_ids": [40],
        "path": WHATSAPP_EVIDENCE_DIR / "ticket_40.png",
        "description": "Mensaje inicial enviado a WhatsApp con ticket #40 y captura compartida por la usuaria.",
    },
]
WHATSAPP_EVIDENCE_TICKET_IDS = sorted({ticket_id for item in WHATSAPP_EVIDENCE for ticket_id in item["ticket_ids"]})

GREEN = "16A34A"
DARK_GREEN = "006837"
PURPLE = "7E22CE"
DARK = "1F2937"
GRAY = "6B7280"
LIGHT_GREEN = "DCFCE7"
LIGHT_PURPLE = "F3E8FF"
LIGHT_BLUE = "DBEAFE"
LIGHT_AMBER = "FEF3C7"

STATUS_LABEL = {
    "open": "Abierto",
    "closed": "Cerrado",
    "escalated_whatsapp": "Escalado a WhatsApp",
    None: "Sin estado",
}
CHANNEL_LABEL = {
    "whatsapp": "WhatsApp",
    "ia": "IA",
    None: "Sin canal final",
    "": "Sin canal final",
}
TOPIC_LABEL = {
    "manual usuario tecnico estudiantes": "Manual usuario técnico estudiantes",
    "saludo": "Saludo",
    "progreso": "Progreso",
    "acceso_login": "Acceso / login",
    "modulos": "Módulos",
    "dashboard": "Dashboard",
    "contacto_humano": "Contacto humano",
    "plan_negocio_tecnico": "Plan de negocio técnico",
    "sincronizacion": "Sincronización",
    "desbloqueo": "Desbloqueo",
    "perfil": "Perfil",
    None: "Sin tópico",
    "": "Sin tópico",
}
SCREEN_LABEL = {
    "student-node-forum": "Foro del estudiante",
    "student-dashboard": "Dashboard estudiante",
    None: "Sin pantalla registrada",
    "": "Sin pantalla registrada",
}


def aws_json(args):
    raw = subprocess.check_output(args, text=True, encoding="utf-8")
    return json.loads(raw)


def aws_text(args):
    return subprocess.check_output(args, text=True, encoding="utf-8")


def list_keys(prefix):
    args = [
        "aws", "s3api", "list-objects-v2",
        "--bucket", BUCKET,
        "--prefix", prefix,
        "--query", "Contents[].Key",
        "--output", "json",
    ]
    keys = aws_json(args) or []
    return [k for k in keys if k and k.endswith(".json")]


def read_s3_json(key):
    raw = aws_text(["aws", "s3", "cp", f"s3://{BUCKET}/{key}", "-"])
    return json.loads(raw)


def parse_dt(value):
    if not value:
        return None
    try:
        return datetime.fromisoformat(value)
    except Exception:
        return None


def fmt_dt(value):
    dt = parse_dt(value)
    return dt.strftime("%Y-%m-%d %H:%M") if dt else ""


def label(mapping, value):
    return mapping.get(value, value or mapping.get(None, "Sin dato"))


def safe_text(value):
    if value is None:
        return ""
    return str(value).replace("\r", " ").replace("\n", " ").strip()


def collect_data():
    latest_keys = list_keys(LATEST_PREFIX)
    history_keys = list_keys(HISTORY_PREFIX)
    records = []
    messages = []
    for key in latest_keys:
        payload = read_s3_json(key)
        ticket = payload.get("ticket") or {}
        ticket_messages = ticket.get("messages") or []
        satisfaction = ticket.get("satisfaction") or {}
        assistant_confidences = [
            m.get("confidence") for m in ticket_messages
            if m.get("sender") == "assistant" and m.get("confidence") is not None
        ]
        has_manual_escalation = any(
            (m.get("sender") == "system" and "whatsapp" in (m.get("message") or "").lower())
            for m in ticket_messages
        )
        has_assistant_whatsapp = any(
            (m.get("sender") == "assistant" and "whatsapp" in (m.get("message") or "").lower())
            for m in ticket_messages
        )
        channel_final = ticket.get("channel_final")
        status = ticket.get("status")
        escalated = channel_final == "whatsapp" or status == "escalated_whatsapp" or has_manual_escalation or has_assistant_whatsapp
        escalation_type = ""
        if escalated:
            if has_manual_escalation:
                escalation_type = "Manual desde widget"
            elif has_assistant_whatsapp:
                escalation_type = "Automático o sugerido por IA"
            else:
                escalation_type = "Marcado como WhatsApp"

        rec = {
            "key": key,
            "archive_reason": payload.get("archive_reason"),
            "archived_at": payload.get("archived_at"),
            "ticket_id": ticket.get("id"),
            "user_id": ticket.get("user_id"),
            "status": status,
            "status_label": label(STATUS_LABEL, status),
            "topic": ticket.get("topic"),
            "topic_label": label(TOPIC_LABEL, ticket.get("topic")),
            "summary": ticket.get("summary"),
            "resolved": ticket.get("resolved"),
            "channel_final": channel_final,
            "channel_label": label(CHANNEL_LABEL, channel_final),
            "created_at": ticket.get("created_at"),
            "updated_at": ticket.get("updated_at"),
            "screen_label": (payload.get("extra") or {}).get("screen_label"),
            "screen_label_display": label(SCREEN_LABEL, (payload.get("extra") or {}).get("screen_label")),
            "messages_count": len(ticket_messages),
            "user_messages": sum(1 for m in ticket_messages if m.get("sender") == "user"),
            "assistant_messages": sum(1 for m in ticket_messages if m.get("sender") == "assistant"),
            "system_messages": sum(1 for m in ticket_messages if m.get("sender") == "system"),
            "assistant_confidence_avg": round(statistics.mean(assistant_confidences), 3) if assistant_confidences else None,
            "escalated_to_whatsapp": escalated,
            "escalation_type": escalation_type,
            "satisfaction_resolved": satisfaction.get("resolved") if satisfaction else None,
            "satisfaction_rating": satisfaction.get("rating") if satisfaction else None,
            "satisfaction_comment": satisfaction.get("comment") if satisfaction else None,
        }
        records.append(rec)
        for msg in ticket_messages:
            messages.append({
                "ticket_id": ticket.get("id"),
                "message_id": msg.get("id"),
                "sender": msg.get("sender"),
                "message": msg.get("message"),
                "confidence": msg.get("confidence"),
                "created_at": msg.get("created_at"),
                "topic": ticket.get("topic"),
                "status": status,
                "channel_final": channel_final,
            })
    records.sort(key=lambda r: (r.get("created_at") or "", r.get("ticket_id") or 0))
    messages.sort(key=lambda r: (r.get("created_at") or "", r.get("ticket_id") or 0, r.get("message_id") or 0))
    return records, messages, history_keys


def build_summary(records, messages, history_keys):
    by_status = collections.Counter(r["status_label"] for r in records)
    by_channel = collections.Counter(r["channel_label"] for r in records)
    by_topic = collections.Counter(r["topic_label"] for r in records)
    by_screen = collections.Counter(r["screen_label_display"] for r in records)
    by_month_day = collections.Counter((r.get("created_at") or "")[:10] for r in records)
    by_month_day.pop("", None)
    satisfactions = [r for r in records if r.get("satisfaction_rating") is not None or r.get("satisfaction_resolved") is not None]
    ratings = [r["satisfaction_rating"] for r in records if r.get("satisfaction_rating") is not None]
    confidence_values = [m["confidence"] for m in messages if m.get("sender") == "assistant" and m.get("confidence") is not None]
    escalated = [r for r in records if r["escalated_to_whatsapp"]]
    closed = [r for r in records if r["status"] == "closed"]
    open_tickets = [r for r in records if r["status"] == "open"]
    date_values = [r["created_at"] for r in records if r.get("created_at")]
    summary = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "s3_bucket": BUCKET,
        "latest_snapshot_count": len(records),
        "history_snapshot_count": len(history_keys),
        "date_min": min(date_values) if date_values else None,
        "date_max": max(date_values) if date_values else None,
        "tickets_total": len(records),
        "messages_total": len(messages),
        "user_messages": sum(1 for m in messages if m.get("sender") == "user"),
        "assistant_messages": sum(1 for m in messages if m.get("sender") == "assistant"),
        "system_messages": sum(1 for m in messages if m.get("sender") == "system"),
        "closed_total": len(closed),
        "open_total": len(open_tickets),
        "escalated_whatsapp_total": len(escalated),
        "satisfaction_total": len(satisfactions),
        "satisfaction_resolved_true": sum(1 for r in satisfactions if r.get("satisfaction_resolved") is True),
        "satisfaction_avg_rating": round(statistics.mean(ratings), 2) if ratings else None,
        "confidence_avg": round(statistics.mean(confidence_values), 3) if confidence_values else None,
        "confidence_min": min(confidence_values) if confidence_values else None,
        "confidence_max": max(confidence_values) if confidence_values else None,
        "by_status": dict(by_status),
        "by_channel": dict(by_channel),
        "by_topic": dict(by_topic.most_common()),
        "by_screen": dict(by_screen.most_common()),
        "by_day": dict(sorted(by_month_day.items())),
        "latest_records": sorted(records, key=lambda r: r.get("updated_at") or "", reverse=True)[:10],
    }
    return summary


def write_csv(records, messages):
    ticket_fields = [
        "ticket_id", "user_id", "created_at", "updated_at", "status_label", "topic_label", "summary",
        "channel_label", "escalated_to_whatsapp", "escalation_type", "resolved", "messages_count",
        "user_messages", "assistant_messages", "system_messages", "assistant_confidence_avg",
        "screen_label_display", "satisfaction_resolved", "satisfaction_rating", "satisfaction_comment", "key",
    ]
    with MASTER_CSV_OUT.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=ticket_fields)
        writer.writeheader()
        for r in records:
            writer.writerow({k: r.get(k) for k in ticket_fields})

    message_fields = ["ticket_id", "message_id", "created_at", "sender", "message", "confidence", "topic", "status", "channel_final"]
    with MESSAGES_CSV_OUT.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=message_fields)
        writer.writeheader()
        for m in messages:
            writer.writerow({k: m.get(k) for k in message_fields})


def write_xlsx(records, messages, summary):
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws.append(["Indicador", "Valor"])
    rows = [
        ("Tickets analizados", summary["tickets_total"]),
        ("Mensajes registrados", summary["messages_total"]),
        ("Mensajes de usuarios", summary["user_messages"]),
        ("Respuestas del asistente", summary["assistant_messages"]),
        ("Mensajes de sistema", summary["system_messages"]),
        ("Tickets abiertos", summary["open_total"]),
        ("Tickets cerrados", summary["closed_total"]),
        ("Casos con canal WhatsApp o escalamiento", summary["escalated_whatsapp_total"]),
        ("Encuestas registradas", summary["satisfaction_total"]),
        ("Casos marcados como resueltos en encuesta", summary["satisfaction_resolved_true"]),
        ("Calificación promedio", summary["satisfaction_avg_rating"]),
        ("Confianza promedio respuestas IA", summary["confidence_avg"]),
        ("Snapshots históricos en S3", summary["history_snapshot_count"]),
        ("Rango de fechas", f"{summary['date_min']} a {summary['date_max']}"),
    ]
    for row in rows:
        ws.append(row)

    def add_counter_sheet(title, counter_dict):
        sheet = wb.create_sheet(title)
        sheet.append(["Categoría", "Cantidad"])
        for k, v in counter_dict.items():
            sheet.append([k, v])
        return sheet

    add_counter_sheet("Estados", summary["by_status"])
    add_counter_sheet("Canales", summary["by_channel"])
    add_counter_sheet("Temas", summary["by_topic"])
    add_counter_sheet("Pantallas", summary["by_screen"])

    tickets_sheet = wb.create_sheet("Tickets")
    ticket_headers = [
        "ticket_id", "user_id", "created_at", "updated_at", "status_label", "topic_label", "summary",
        "channel_label", "escalated_to_whatsapp", "escalation_type", "resolved", "messages_count",
        "user_messages", "assistant_messages", "system_messages", "assistant_confidence_avg",
        "screen_label_display", "satisfaction_resolved", "satisfaction_rating", "satisfaction_comment", "key",
    ]
    tickets_sheet.append(ticket_headers)
    for r in records:
        tickets_sheet.append([r.get(h) for h in ticket_headers])

    messages_sheet = wb.create_sheet("Mensajes")
    message_headers = ["ticket_id", "message_id", "created_at", "sender", "message", "confidence", "topic", "status", "channel_final"]
    messages_sheet.append(message_headers)
    for m in messages:
        messages_sheet.append([m.get(h) for h in message_headers])

    header_fill = PatternFill("solid", fgColor="7E22CE")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D1D5DB")
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in sheet.iter_rows():
            for cell in row:
                cell.border = Border(top=thin, left=thin, right=thin, bottom=thin)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col in range(1, min(sheet.max_column, 12) + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 22
        if sheet.max_column >= 7:
            sheet.column_dimensions[get_column_letter(7)].width = 55
        if sheet.title == "Mensajes":
            sheet.column_dimensions["E"].width = 80
    wb.save(XLSX_OUT)


def plot_bar(counter_dict, title, filename, color=GREEN, horizontal=False, top_n=None):
    items = list(counter_dict.items())
    if top_n:
        items = items[:top_n]
    if not items:
        return None
    labels = [str(k) for k, _ in items]
    values = [v for _, v in items]
    plt.style.use("default")
    if horizontal:
        fig, ax = plt.subplots(figsize=(9, 5), dpi=180)
        ax.barh(labels[::-1], values[::-1], color=color)
        ax.set_xlabel("Cantidad")
        for idx, val in enumerate(values[::-1]):
            ax.text(val + 0.2, idx, str(val), va="center", fontsize=8)
    else:
        fig, ax = plt.subplots(figsize=(8, 4.6), dpi=180)
        ax.bar(labels, values, color=color)
        ax.set_ylabel("Cantidad")
        for idx, val in enumerate(values):
            ax.text(idx, val + 0.2, str(val), ha="center", fontsize=8)
        ax.tick_params(axis="x", rotation=25)
    ax.set_title(title, fontsize=13, weight="bold")
    ax.grid(axis="y" if not horizontal else "x", alpha=0.25)
    fig.tight_layout()
    path = OUT / filename
    fig.savefig(path, bbox_inches="tight")
    plt.close(fig)
    return path


def plot_charts(summary):
    charts = {}
    charts["status"] = plot_bar(summary["by_status"], "Tickets por estado", "tickets_por_estado.png", color="#7E22CE")
    charts["channel"] = plot_bar(summary["by_channel"], "Tickets por canal final", "tickets_por_canal.png", color="#16A34A")
    charts["topic"] = plot_bar(summary["by_topic"], "Temas más frecuentes", "tickets_por_tema.png", color="#2563EB", horizontal=True, top_n=10)
    charts["screen"] = plot_bar(summary["by_screen"], "Pantalla de origen", "tickets_por_pantalla.png", color="#F59E0B")
    if summary["by_day"]:
        labels = list(summary["by_day"].keys())
        values = list(summary["by_day"].values())
        fig, ax = plt.subplots(figsize=(9, 4.6), dpi=180)
        ax.plot(labels, values, marker="o", color="#7E22CE")
        ax.set_title("Evolución de tickets por fecha de creación", fontsize=13, weight="bold")
        ax.set_ylabel("Tickets")
        ax.grid(axis="y", alpha=0.25)
        ax.tick_params(axis="x", rotation=45)
        for x, y in zip(labels, values):
            ax.text(x, y + 0.05, str(y), ha="center", fontsize=8)
        fig.tight_layout()
        path = OUT / "tickets_por_dia.png"
        fig.savefig(path, bbox_inches="tight")
        plt.close(fig)
        charts["day"] = path
    return charts


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, color=None, size=8.5):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run("" if text is None else str(text))
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = RGBColor.from_string(DARK if level == 1 else PURPLE)
    return p


def add_table(doc, headers, rows, header_fill=PURPLE, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, bold=True, color="FFFFFF", size=font_size)
        set_cell_shading(table.rows[0].cells[i], header_fill)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val, size=font_size)
    doc.add_paragraph()
    return table


def add_metric_cards(doc, metrics):
    table = doc.add_table(rows=2, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, (title, value) in enumerate(metrics[:4]):
        cell = table.rows[0].cells[idx]
        set_cell_text(cell, title, bold=True, color="FFFFFF", size=8)
        set_cell_shading(cell, DARK_GREEN)
        cell = table.rows[1].cells[idx]
        set_cell_text(cell, value, bold=True, color=DARK, size=12)
        set_cell_shading(cell, LIGHT_GREEN)
    doc.add_paragraph()


def add_picture_if_exists(doc, path, width=6.2):
    if path and Path(path).exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))


def write_docx(records, messages, history_keys, summary, charts):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("INFORME DE INTERACCIONES DEL CENTRO DE AYUDA")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = RGBColor.from_string(DARK)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Nodo Cordillera — Plataforma EmprendiPaz")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string(PURPLE)
    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_p.add_run(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}").italic = True

    add_heading(doc, "1. Propósito y alcance", 1)
    doc.add_paragraph(
        "Este informe presenta el estado de las interacciones registradas en el centro de ayuda de la plataforma EmprendiPaz, "
        "con énfasis en la atención inicial por asistente de IA, los casos escalados a WhatsApp y la retroalimentación de satisfacción registrada. "
        "El análisis se prepara para la presentación del Nodo Cordillera, integrado por los municipios de Cumbitara, El Rosario, Leiva, Policarpa y Taminango."
    )
    doc.add_paragraph(
        "La fuente primaria de evidencia son los snapshots JSON archivados en S3 bajo support/conversations/latest y support/conversations/history. "
        "Estos snapshots contienen el ticket, mensajes registrados, estado, canal final y satisfacción cuando fue reportada."
    )

    add_heading(doc, "2. Resumen ejecutivo", 1)
    add_metric_cards(doc, [
        ("Tickets analizados", str(summary["tickets_total"])),
        ("Mensajes registrados", str(summary["messages_total"])),
        ("Casos WhatsApp", str(summary["escalated_whatsapp_total"])),
        ("Satisfacción promedio", str(summary["satisfaction_avg_rating"] or "Sin dato")),
    ])
    doc.add_paragraph(
        f"Se analizaron {summary['tickets_total']} tickets y {summary['messages_total']} mensajes registrados. "
        f"Del total, {summary['assistant_messages']} corresponden a respuestas del asistente, {summary['user_messages']} a mensajes de usuarios y {summary['system_messages']} a eventos del sistema. "
        f"Se identificaron {summary['escalated_whatsapp_total']} casos con canal final o escalamiento a WhatsApp, {summary['closed_total']} tickets cerrados y {summary['open_total']} tickets abiertos."
    )
    doc.add_paragraph(
        f"La satisfacción fue reportada en {summary['satisfaction_total']} tickets; {summary['satisfaction_resolved_true']} fueron marcados como resueltos y la calificación promedio fue {summary['satisfaction_avg_rating']}/5. "
        f"La confianza promedio registrada en respuestas del asistente fue {summary['confidence_avg']}, con valores entre {summary['confidence_min']} y {summary['confidence_max']}."
    )

    add_heading(doc, "3. Metodología y trazabilidad", 1)
    add_table(doc, ["Elemento", "Descripción"], [
        ["Fuente S3", f"s3://{BUCKET}/support/conversations/latest/ y history/"],
        ["Snapshots latest", summary["latest_snapshot_count"]],
        ["Snapshots históricos", summary["history_snapshot_count"]],
        ["Rango de creación", f"{fmt_dt(summary['date_min'])} a {fmt_dt(summary['date_max'])}"],
        ["Evidencia exportada", f"{XLSX_OUT.name} y CSV en {OUT}"],
    ], font_size=8.5)
    doc.add_paragraph(
        "La consolidación utiliza el snapshot latest de cada ticket como estado más reciente y conserva el historial de snapshots como evidencia de trazabilidad del ciclo de vida. "
        "La conversación posterior realizada directamente en WhatsApp no está integrada por API en la plataforma; por tanto, el informe registra el evento de escalamiento y el canal final, no el contenido del chat humano externo."
    )

    add_heading(doc, "4. Resultados cuantitativos", 1)
    add_heading(doc, "4.1 Tickets por estado", 2)
    add_table(doc, ["Estado", "Cantidad"], list(summary["by_status"].items()))
    add_picture_if_exists(doc, charts.get("status"))

    add_heading(doc, "4.2 Tickets por canal final", 2)
    add_table(doc, ["Canal", "Cantidad"], list(summary["by_channel"].items()))
    add_picture_if_exists(doc, charts.get("channel"))

    add_heading(doc, "4.3 Temas más frecuentes", 2)
    add_table(doc, ["Tema", "Cantidad"], list(summary["by_topic"].items())[:10])
    add_picture_if_exists(doc, charts.get("topic"))

    add_heading(doc, "4.4 Pantalla de origen", 2)
    add_table(doc, ["Pantalla", "Cantidad"], list(summary["by_screen"].items()))
    add_picture_if_exists(doc, charts.get("screen"))

    add_heading(doc, "4.5 Evolución temporal", 2)
    add_table(doc, ["Fecha", "Tickets creados"], list(summary["by_day"].items()))
    add_picture_if_exists(doc, charts.get("day"))

    add_heading(doc, "5. Análisis de interacción IA", 1)
    doc.add_paragraph(
        "El asistente de IA funciona como primera línea de soporte operativo: recibe el mensaje inicial, recupera contexto desde la base de conocimiento y documentos de soporte, "
        "y genera una respuesta acotada al material disponible. Cada respuesta queda persistida como mensaje del asistente con metadatos de confianza cuando aplican."
    )
    doc.add_paragraph(
        "Los temas más frecuentes se concentran en uso técnico de la plataforma, foro, progreso, acceso/login, módulos y dashboard. "
        "Esto indica que el centro de ayuda está siendo usado principalmente para resolver dudas de navegación, registro de avance y uso de espacios de interacción del estudiante."
    )

    add_heading(doc, "6. Escalamiento a WhatsApp", 1)
    manual_count = sum(1 for r in records if r.get("escalation_type") == "Manual desde widget")
    auto_count = sum(1 for r in records if r.get("escalation_type") == "Automático o sugerido por IA")
    marked_count = sum(1 for r in records if r.get("escalation_type") == "Marcado como WhatsApp")
    add_table(doc, ["Tipo de escalamiento detectado", "Cantidad"], [
        ["Manual desde widget", manual_count],
        ["Automático o sugerido por IA", auto_count],
        ["Marcado como WhatsApp", marked_count],
        ["Total con evidencia de WhatsApp", summary["escalated_whatsapp_total"]],
    ])
    doc.add_paragraph(
        "La plataforma registra el momento en que un caso se escala a WhatsApp mediante estado del ticket, canal final y mensajes de sistema o del asistente. "
        "El enlace generado usa wa.me con un texto prellenado que incluye el id del ticket, el usuario y el resumen del caso, facilitando trazabilidad manual."
    )
    warning = doc.add_paragraph()
    run = warning.add_run("Nota de auditoría: ")
    run.bold = True
    warning.add_run(
        "no se encontró integración bidireccional con WhatsApp, Twilio, Meta Webhooks o WhatsApp Business API. "
        "Por ello, el contenido posterior de la conversación humana en WhatsApp no queda registrado dentro de la plataforma. Para auditar ese tramo se requiere una exportación manual del WhatsApp de soporte."
    )

    add_heading(doc, "7. Satisfacción y cierre", 1)
    add_table(doc, ["Indicador", "Resultado"], [
        ["Encuestas registradas", summary["satisfaction_total"]],
        ["Marcadas como resueltas", summary["satisfaction_resolved_true"]],
        ["Calificación promedio", summary["satisfaction_avg_rating"] or "Sin dato"],
        ["Tickets cerrados", summary["closed_total"]],
    ])
    doc.add_paragraph(
        "La muestra de satisfacción es limitada frente al total de tickets; se recomienda incentivar el cierre formal para mejorar la lectura de efectividad real del centro de ayuda."
    )

    add_heading(doc, "8. Recomendaciones", 1)
    recommendations = [
        "Crear una vista o exportación administrativa para descargar tickets, mensajes y satisfacción sin depender manualmente de S3.",
        "Incentivar que estudiantes cierren el caso con encuesta para mejorar métricas de resolución y calidad.",
        "Etiquetar explícitamente los tickets escalados por baja confianza, por solicitud humana y por botón manual para separar causas de escalamiento.",
        "Si el seguimiento por WhatsApp debe auditarse, integrar WhatsApp Business API o definir un procedimiento de exportación periódica del chat de soporte.",
        "Actualizar la base de conocimiento con los temas frecuentes detectados: foro, progreso, módulos, dashboard, acceso y perfil.",
    ]
    for rec in recommendations:
        doc.add_paragraph(rec, style="List Bullet")

    add_heading(doc, "9. Anexo: últimos tickets actualizados", 1)
    latest_rows = []
    for r in summary["latest_records"]:
        latest_rows.append([
            r.get("ticket_id"),
            fmt_dt(r.get("updated_at")),
            r.get("status_label"),
            r.get("topic_label"),
            r.get("channel_label"),
            safe_text(r.get("summary"))[:120],
        ])
    add_table(doc, ["Ticket", "Última actualización", "Estado", "Tema", "Canal", "Resumen"], latest_rows, font_size=7.5)

    add_heading(doc, "10. Anexo: evidencias WhatsApp recibidas", 1)
    whatsapp_records = [r for r in records if r.get("escalated_to_whatsapp")]
    doc.add_paragraph(
        "El tutor compartió capturas puntuales de conversaciones de WhatsApp asociadas a algunos tickets escalados desde la plataforma. "
        "Por tanto, el anexo se limita a las evidencias realmente disponibles y no deja espacios pendientes para soportes no entregados."
    )
    doc.add_paragraph(
        f"Tickets con captura externa recibida: {', '.join('#' + str(ticket_id) for ticket_id in WHATSAPP_EVIDENCE_TICKET_IDS)}. "
        f"Los demás casos marcados como WhatsApp se sustentan únicamente con la evidencia interna del ticket: estado, canal final, mensajes de sistema o enlace wa.me generado."
    )
    if whatsapp_records:
        whatsapp_rows = []
        for r in whatsapp_records:
            has_external_capture = r.get("ticket_id") in WHATSAPP_EVIDENCE_TICKET_IDS
            whatsapp_rows.append([
                f"Ticket #{r.get('ticket_id')}",
                fmt_dt(r.get("created_at")),
                r.get("status_label"),
                r.get("topic_label"),
                r.get("escalation_type") or "Escalado a WhatsApp",
                safe_text(r.get("summary"))[:110],
                "Captura recibida" if has_external_capture else "Sin captura externa; evidencia interna de escalamiento",
            ])
        add_table(
            doc,
            ["Ticket", "Fecha", "Estado", "Tema", "Tipo", "Resumen", "Soporte disponible"],
            whatsapp_rows,
            font_size=7,
        )

        records_by_id = {r.get("ticket_id"): r for r in records}
        for item in WHATSAPP_EVIDENCE:
            existing_ticket_ids = [ticket_id for ticket_id in item["ticket_ids"] if ticket_id in records_by_id]
            ticket_label = ", ".join(f"#{ticket_id}" for ticket_id in item["ticket_ids"])
            p = doc.add_paragraph()
            run = p.add_run(f"Evidencia WhatsApp - Ticket(s) {ticket_label}: ")
            run.bold = True
            p.add_run(item["description"])
            for ticket_id in existing_ticket_ids:
                rec = records_by_id[ticket_id]
                doc.add_paragraph(
                    f"Resumen interno ticket #{ticket_id}: {safe_text(rec.get('summary')) or 'Sin resumen registrado.'}",
                    style="List Bullet",
                )
            if item["path"].exists():
                add_picture_if_exists(doc, item["path"], width=3.15)
            else:
                doc.add_paragraph(f"No se encontró el archivo local de la captura: {item['path']}")
            doc.add_paragraph()
    else:
        doc.add_paragraph("No se encontraron tickets con canal final o escalamiento a WhatsApp.")

    add_heading(doc, "11. Referencias técnicas", 1)
    refs = [
        "frontend/frontend-app/src/components/student/SupportCenterWidget.jsx",
        "backend/backend-app/src/routes/support.py",
        "backend/backend-app/src/services/support_service.py",
        "backend/backend-app/src/models/support_ticket.py",
        "docs/CENTRO_AYUDA_ARQUITECTURA_AUDITORIA.md",
        "s3://elearning-archivos/support/conversations/latest/",
        "s3://elearning-archivos/support/conversations/history/",
    ]
    for ref in refs:
        doc.add_paragraph(ref, style="List Bullet")

    doc.save(DOCX_OUT)


def main():
    records, messages, history_keys = collect_data()
    summary = build_summary(records, messages, history_keys)
    SUMMARY_JSON_OUT.write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    write_csv(records, messages)
    write_xlsx(records, messages, summary)
    charts = plot_charts(summary)
    write_docx(records, messages, history_keys, summary, charts)
    print(json.dumps({
        "docx": str(DOCX_OUT),
        "xlsx": str(XLSX_OUT),
        "summary_json": str(SUMMARY_JSON_OUT),
        "tickets_csv": str(MASTER_CSV_OUT),
        "messages_csv": str(MESSAGES_CSV_OUT),
        "charts": {k: str(v) for k, v in charts.items() if v},
        "tickets_total": summary["tickets_total"],
        "messages_total": summary["messages_total"],
        "whatsapp_total": summary["escalated_whatsapp_total"],
    }, ensure_ascii=False, indent=2))

if __name__ == "__main__":
    main()
