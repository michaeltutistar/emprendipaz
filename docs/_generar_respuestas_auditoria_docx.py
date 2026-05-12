# Genera Respuestas_Cuestionario_Auditoria_Aplicaciones_Externas.docx (ejecutar una vez).
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

OUT = Path(__file__).resolve().parent / "Respuestas_Cuestionario_Auditoria_Aplicaciones_Externas.docx"

ITEMS = [
    (
        "Contexto",
        "Alcance de esta respuesta",
        "Este documento responde al cuestionario de auditoría de aplicaciones externas respecto de la plataforma "
        "EmprendiPaz / e-learning (frontend en https://emprendimiento-narino.com, API en AWS Lambda, base PostgreSQL en RDS, "
        "archivos en S3), según lo implementado en el repositorio técnico y la configuración documentada en DEPLOYMENT.md. "
        "Los aspectos puramente organizativos o legales (p. ej. política de privacidad publicada como texto jurídico final) "
        "deben complementarse con documentación institucional de la Gobernación.",
    ),
    (
        "1. Seguridad y control de acceso — Autenticación",
        "¿El sistema implementa autenticación segura (hash de contraseñas — bcrypt, Argon2)?",
        "Las contraseñas no se almacenan en texto plano. Se utiliza Werkzeug (generate_password_hash / check_password_hash) "
        "sobre el modelo de usuario, esquema de derivación de clave estándar de Werkzeug (pbkdf2:sha256 en las versiones habituales). "
        "No se usa bcrypt ni Argon2 de forma explícita; sí hay hash unidireccional y verificación segura al iniciar sesión.",
    ),
    (
        "1. Seguridad y control de acceso — Autenticación",
        "¿Existe protección contra ataques de fuerza bruta?",
        "No hay limitación de intentos de login ni bloqueo progresivo implementado en el código de la API revisado. "
        "Es posible complementar con cuotas/throttling en API Gateway, WAF o herramientas de borde en AWS; se recomienda "
        "explicitar en gobierno de infraestructura si aplica.",
    ),
    (
        "1. Seguridad y control de acceso — Autenticación",
        "¿Se utiliza CAPTCHA o verificación adicional?",
        "No se identifica CAPTCHA u otro desafío en el flujo de login/registro en el código actual.",
    ),
    (
        "1. Seguridad y control de acceso — Autenticación",
        "¿Se maneja doble factor de autenticación (2FA)?",
        "No está implementado 2FA/MFA en la aplicación actual.",
    ),
    (
        "1. Seguridad y control de acceso — Autorización",
        "¿Existen roles definidos? (usuario, evaluador, admin, etc.)",
        "Sí. El backend distingue al menos los roles: administrador, evaluador, instructor y estudiante, con decoradores "
        "de autorización (p. ej. admin_required, evaluador_required, instructor_required, student_required) en auth_service.",
    ),
    (
        "1. Seguridad y control de acceso — Autorización",
        "¿Los permisos están correctamente segregados?",
        "La segregación se aplica por rutas (blueprints /api/admin, /api/instructor, /api/student, etc.) y decoradores que "
        "verifican el rol del usuario autenticado antes de ejecutar la lógica. Las operaciones sensibles exigen JWT válido.",
    ),
    (
        "1. Seguridad y control de acceso — Autorización",
        "¿Se puede acceder a módulos sin autenticación?",
        "El frontend puede mostrar páginas públicas (p. ej. aterrizaje, información). Las APIs de datos personales, "
        "progreso, administración e instructor requieren token (Authorization: Bearer) salvo endpoints explícitamente "
        "públicos como login/registro/recuperación según diseño.",
    ),
    (
        "1. Seguridad y control de acceso — Sesiones",
        "¿Las sesiones expiran correctamente?",
        "El access token JWT tiene caducidad configurada (24 horas en generate_token). El refresh token tiene vigencia "
        "de 30 días y se entrega en cookie HttpOnly. Los tokens expirados son rechazados al decodificar.",
    ),
    (
        "1. Seguridad y control de acceso — Sesiones",
        "¿Se regeneran los tokens de sesión?",
        "Cada inicio de sesión exitoso genera nuevos access y refresh tokens. La renovación mediante refresh está contemplada "
        "en el servicio de autenticación (verify_refresh_token).",
    ),
    (
        "1. Seguridad y control de acceso — Sesiones",
        "¿Se protege contra secuestro de sesión?",
        "Se usa HTTPS en producción (CloudFront con redirección a HTTPS en la configuración del distribuidor). La cookie de "
        "refresh se marca Secure y SameSite acorde al entorno; el access token viaja preferentemente en cabecera Authorization, "
        "reduciendo exposición a CSRF clásico en cookies de sesión.",
    ),
    (
        "1. Seguridad y control de acceso — Seguridad web",
        "¿Protección contra SQL Injection?",
        "Las consultas a base de datos se realizan principalmente mediante SQLAlchemy ORM con parámetros enlazados, "
        "lo que mitiga inyección SQL en el acceso habitual a datos.",
    ),
    (
        "1. Seguridad y control de acceso — Seguridad web",
        "¿Protección contra XSS?",
        "El frontend es React; el renderizado por defecto escapa contenido. Cualquier uso futuro de HTML crudo debería "
        "mantener saneamiento explícito.",
    ),
    (
        "1. Seguridad y control de acceso — Seguridad web",
        "¿Protección contra CSRF?",
        "La API es de estilo REST con JWT en cabecera para la mayoría de operaciones; la cookie de refresh tiene atributos "
        "SameSite. El riesgo CSRF típico de formularios con cookie de sesión única es menor que en aplicaciones solo-cookie.",
    ),
    (
        "1. Seguridad y control de acceso — Seguridad web",
        "¿Se usan cabeceras de seguridad (CSP, HSTS, etc.)?",
        "Hay redirección forzada a HTTPS en CloudFront para el sitio estático. Políticas explícitas CSP o HSTS a nivel de "
        "Response Headers Policy no están detalladas en el repositorio; pueden añadirse en CloudFront/API Gateway para endurecer el perfil.",
    ),
    (
        "2. Arquitectura y tecnología",
        "¿Qué stack tecnológico utiliza?",
        "Frontend: JavaScript/React (Vite). Backend: Python 3, Flask. Base de datos: PostgreSQL (RDS). "
        "Almacenamiento de objetos: Amazon S3. Cómputo API: AWS Lambda + API Gateway. Entrega web: S3 + CloudFront.",
    ),
    (
        "2. Arquitectura y tecnología",
        "¿La arquitectura es monolítica, modular o basada en servicios (API REST)?",
        "Backend modular por dominios (blueprints) expuesto como API REST; despliegue como una función Lambda que concentra la app. "
        "No es un monolito desplegado en VM única tradicional, sino serverless con capa de API.",
    ),
    (
        "2. Arquitectura y tecnología",
        "¿Existe separación entre frontend, backend y base de datos?",
        "Sí. Frontend estático independiente; backend sin UI embebida; persistencia en RDS y archivos en S3.",
    ),
    (
        "2. Arquitectura y tecnología",
        "¿Se implementa control de versiones (Git)?",
        "El proyecto se mantiene en repositorio Git (ramas, historial de cambios).",
    ),
    (
        "2. Arquitectura y tecnología",
        "¿Existe documentación técnica?",
        "Existe documentación de despliegue y componentes en DEPLOYMENT.md (S3, CloudFront, Lambda, empaquetado ZIP, variables).",
    ),
    (
        "3. Base de datos y gestión de información",
        "¿El modelo de datos está normalizado?",
        "El modelo relacional con SQLAlchemy incluye entidades separadas (usuarios, cursos, módulos, inscripciones, foros, etc.) "
        "con relaciones, en línea con normalización práctica para la aplicación.",
    ),
    (
        "3. Base de datos y gestión de información",
        "¿Se utilizan índices correctamente?",
        "Varias tablas definen índices y claves foráneas (p. ej. foro por nodo/autor) para consultas frecuentes.",
    ),
    (
        "3. Base de datos y gestión de información",
        "¿Hay integridad referencial (FK)?",
        "Sí, mediante ForeignKey en modelos SQLAlchemy (p. ej. usuarios, hilos y respuestas de foro, tickets de soporte, planes de negocio).",
    ),
    (
        "3. Base de datos y gestión de información",
        "¿Se realizan backups automáticos?",
        "Los backups automáticos y ventanas de retención son responsabilidad de la configuración operativa de Amazon RDS "
        "(servicio gestionado). Conviene adjuntar a la auditoría la política de backup/restore acordada con TI.",
    ),
    (
        "3. Base de datos y gestión de información",
        "¿Existe control de duplicidad de registros?",
        "Hay validaciones de negocio (p. ej. email de usuario) y reglas en modelo de datos para evitar duplicados lógicos relevantes.",
    ),
    (
        "3. Base de datos y gestión de información",
        "¿Se manejan logs de auditoría (quién hizo qué)?",
        "Existe la tabla/modelo LogActividad con usuario, acción, detalles y fecha, usada para registrar actividades "
        "(p. ej. avance formativo). Complementariamente, acciones administrativas pueden dejar rastro en la BD según el flujo.",
    ),
    (
        "3. Base de datos y gestión de información",
        "¿Motor gestor de la BD PostgreSQL?",
        "Sí, PostgreSQL como motor indicado en el stack (RDS).",
    ),
    (
        "4. Rendimiento y escalabilidad",
        "¿Cuál es el tiempo de carga promedio?",
        "No hay métrica fija en el repositorio; el frontend se sirve por CloudFront con compresión habilitada en la distribución, "
        "lo que reduce latencia percibida. Se puede aportar informe de monitoreo (CloudWatch/RUM) si TI lo tiene desplegado.",
    ),
    (
        "4. Rendimiento y escalabilidad",
        "¿Soporta múltiples usuarios concurrentes? (ej.: 728+)",
        "Lambda escala concurrentemente según límites de cuenta y configuración; RDS dimensiona según instancia elegida. "
        "La arquitectura está orientada a multiusuario; la cifra concreta de usuarios simultáneos depende de capacidad "
        "provisionada y pruebas de carga.",
    ),
    (
        "4. Rendimiento y escalabilidad",
        "¿Se usa caché?",
        "Caché en borde: CloudFront para activos estáticos del frontend. Caché de aplicación explícita (Redis, etc.) no aparece "
        "como requisito central en el código revisado.",
    ),
    (
        "4. Rendimiento y escalabilidad",
        "¿Se optimizan consultas SQL?",
        "Se usan consultas ORM acotadas; optimizaciones puntuales (índices, filtros) según modelos. Para auditoría de rendimiento "
        "se puede añadir revisión de planes de ejecución en RDS.",
    ),
    (
        "4. Rendimiento y escalabilidad",
        "¿El sistema soporta crecimiento futuro?",
        "Sí: escalado horizontal de invocaciones Lambda, ampliación de RDS, y almacenamiento ilimitado práctico en S3.",
    ),
    (
        "5. Experiencia de usuario (UX/UI)",
        "Navegación: menú claro; registro, login, módulos, estado del proceso",
        "La aplicación separa flujos por rol (estudiante, instructor, admin) con paneles dedicados. El registro es multi-paso "
        "con validaciones y carga de documentos. El login entrega token y perfil de usuario.",
    ),
    (
        "5. Experiencia de usuario (UX/UI)",
        "Diseño responsive, accesible (WCAG), consistencia visual",
        "Interfaz React adaptable a distintos tamaños de pantalla en uso corriente. No hay en el repositorio un informe de "
        "certificación WCAG; se puede declarar cumplimiento parcial y roadmap de mejora de accesibilidad si la entidad lo exige.",
    ),
    (
        "5. Experiencia de usuario (UX/UI)",
        "Flujo de usuario: registro, guía paso a paso, mensajes de error",
        "Registro guiado por pasos con validación en cliente y servidor; mensajes de error JSON en API y manejo en UI.",
    ),
    (
        "6. Módulos de formación y contenido",
        "Estructura pedagógica, modelo ADDIE, progreso, certificación, evaluación",
        "Hay cursos, módulos y lecciones con seguimiento de progreso y actividades (incl. evaluaciones y talleres en distintas "
        "unidades). El programa cubre líneas temáticas del proyecto (liderazgo, finanzas, marketing, modelos de negocio). "
        "Certificación formal depende del marco del programa institucional más que del software aislado.",
    ),
    (
        "7. Formularios y gestión documental",
        "Validación, tamaño y tipo de archivos, subida, subsanación, historial, trazabilidad",
        "Los endpoints de carga (p. ej. file_upload) validan extensión contra una lista permitida y tamaño máximo (10 MB). "
        "Se usa secure_filename y almacenamiento en S3 con metadatos (quién subió). El registro permite completar/información "
        "y documentos según flujo multi-paso. Historial versionado documental granular debe detallarse según política de negocio; "
        "la BD conserva estados y referencias a archivos.",
    ),
    (
        "8. Notificaciones y comunicación",
        "Notificaciones de estado, errores, aprobación/rechazo; correo; bandeja interna",
        "Existe modelo de notificaciones en backend y rutas asociadas. Pueden implementarse avisos en aplicación; correo "
        "transaccional depende de integración SMTP/SES si está configurada en despliegue (no fijada en el fragmento de código base). "
        "Hay módulo de soporte/tickets y asistente con documentación en S3; opcionalmente integración con OpenRouter para respuestas "
        "automáticas (servicio externo bajo variables de entorno), que debe declararse en auditoría de proveedores.",
    ),
    (
        "9. Reportes y analítica",
        "Reportes de usuarios, avance, aprobaciones; dashboards; exportación Excel/PDF",
        "Paneles de administración e instructor incluyen métricas y listados de usuarios/actividad. La exportación a Excel/PDF "
        "depende de funciones específicas desplegadas (p. ej. informes en admin); el administrador puede extraer datos vía "
        "reportes de la plataforma y consultas autorizadas.",
    ),
    (
        "10. Gobernanza y trazabilidad",
        "Registro de quién evalúa, cuándo, qué cambió; historial por usuario; auditoría de decisiones",
        "Estados de cuenta, inscripción y flujos de evaluación quedan en base de datos; LogActividad aporta trazas de acciones "
        "de usuario. Las decisiones de evaluador/admin deben mapearse a tablas/campos concretos del modelo para evidencia en auditoría.",
    ),
    (
        "11. Integración e interoperabilidad",
        "Integración con sistemas de la Gobernación, BD externas, API, escalabilidad a otros programas",
        "La plataforma expone API REST propia consumida por el frontend. No hay en el código una integración nativa obligatoria "
        "con un sistema único de la Gobernación; cualquier integración futura puede hacerse vía API o ETL. El asistente de "
        "soporte puede consumir OpenRouter (IA externa) bajo configuración — conviene política de datos con terceros.",
    ),
    (
        "12. Accesibilidad y multiplataforma",
        "Móvil, tablet, escritorio; baja conectividad; zonas rurales",
        "SPA accesible desde navegadores en distintos dispositivos. El rendimiento en baja conectividad mejora con CDN y "
        "activos estáticos cacheables; no hay modo offline completo declarado en el repositorio.",
    ),
    (
        "13. Cumplimiento normativo",
        "Ley de protección de datos (Habeas Data), políticas de privacidad, términos visibles",
        "En el flujo de registro se exige aceptación de términos y documentos de convocatoria (PDFs en /Terminos y similares). "
        "El texto legal definitivo y el aviso de privacidad institucional deben ser los aprobados por la Gobernación y publicados "
        "en los canales oficiales; el sistema registra flags de aceptación en el usuario.",
    ),
    (
        "14. Calidad del código",
        "Documentación, estructura, buenas prácticas, pruebas (testing)",
        "Código organizado en routes, services, models. No hay una suite de tests de integración prominente en el árbol principal "
        "del producto; se recomienda plan de pruebas automatizadas y revisión de seguridad periódica.",
    ),
    (
        "15. Impacto y efectividad del sistema",
        "Cumplimiento del objetivo, completitud, deserción, indicadores",
        "La plataforma está alineada con la estrategia de formación y acompañamiento a emprendedores en territorio (metas del programa). "
        "Indicadores de completitud y participación pueden obtenerse de la BD (progreso, logs). La deserción y éxito formativo "
        "son analíticas de negocio que la Secretaría puede cuantificar con reportes sobre esos datos.",
    ),
    (
        "Matriz de calificación",
        "Uso de la escala sugerida en el cuestionario",
        "La calificación por criterio (1–5) es competencia del equipo auditor. Este documento solo aporta evidencia técnica "
        "para que cada ítem sea calificado con base en hechos verificables y brechas declaradas explícitamente.",
    ),
]


def main():
    doc = Document()
    title = doc.add_heading(
        "Respuestas al cuestionario de auditoría — Aplicaciones externas",
        0,
    )
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    sub = doc.add_paragraph()
    sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = sub.add_run(
        "Plataforma EmprendiPaz / E-learning — Gobernación de Nariño\n"
        "Documento generado para apoyo a auditoría integral (tecnológica, operativa y estratégica)."
    )
    run.font.size = Pt(11)

    doc.add_paragraph()
    current_section = None
    n = 0
    for section, question, answer in ITEMS:
        if section != current_section:
            current_section = section
            doc.add_heading(section, level=1)
        n += 1
        p_q = doc.add_paragraph()
        p_q.add_run(f"Pregunta / ítem ({n}). ").bold = True
        p_q.add_run(question)
        p_a = doc.add_paragraph()
        p_a.add_run("Respuesta: ").bold = True
        p_a.add_run(answer)
        doc.add_paragraph()

    doc.save(OUT)
    print("Escrito:", OUT)


if __name__ == "__main__":
    main()
