from flask import Blueprint, jsonify, request, session, send_file, make_response  # pyright: ignore[reportMissingImports]
from sqlalchemy import or_
from src.models import db, User, Curso, Inscripcion, LogActividad, CuposConfig, MunicipioCupo, Notificacion
from src.models.intentos_evaluacion import IntentosEvaluacion
from src.models.puntos_plan_negocio import PuntosPlanNegocio
from src.models.respuestas_plan_negocio import RespuestasPlanNegocio
from src.models.asistencia_jornada import AsistenciaJornada
from src.models.formulario_campo import FormularioCampo
from src.models.criterio_evaluacion_config import CriterioEvaluacionConfig
from src.models.cupos_municipio_config import CuposMunicipioConfig
from src.models.documento_config import DocumentoConfig
from src.services.auth_service import token_required, admin_required, admin_or_evaluador_required
from src.services.student_municipio_service import get_preferred_municipio
from datetime import datetime, timedelta
import csv
import io
import tempfile
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

admin_bp = Blueprint('admin', __name__)

# Decorador personalizado eliminado - usando decoradores directos


def _extract_json_payload(flask_response):
    if isinstance(flask_response, tuple):
        response = flask_response[0]
    else:
        response = flask_response
    return response.get_json(silent=True) or {}


def _build_municipios_summary(progreso_estudiantes):
    grupos = {}

    for progreso in progreso_estudiantes or []:
        estudiante_id = progreso.get('estudiante_id')
        estudiante_nombre = progreso.get('estudiante')
        municipio = get_preferred_municipio(estudiante_id, estudiante_nombre, progreso.get('municipio')) or 'Sin municipio'
        if municipio == 'Sin municipio':
            continue

        grupos.setdefault(municipio, []).append(progreso)

    resumen = []
    for municipio, estudiantes in sorted(grupos.items(), key=lambda item: item[0]):
        total_estudiantes = len(estudiantes)
        intentos_totales = 0
        suma_porcentaje_completado = 0
        jornadas_marcadas_total = 0

        for progreso in estudiantes:
            modulos = progreso.get('modulos') or [progreso]
            for modulo_data in modulos:
                progreso_pasos = modulo_data.get('progreso_pasos') or progreso.get('progreso_pasos') or []
                intentos_totales += sum((paso.get('intentos') or 0) for paso in progreso_pasos)

            suma_porcentaje_completado += progreso.get('porcentaje_total') or progreso.get('porcentaje') or 0
            jornadas = progreso.get('jornadas') or {}
            jornadas_marcadas_total += sum(1 for value in jornadas.values() if value)

        pct_completado = round((suma_porcentaje_completado / total_estudiantes), 2) if total_estudiantes > 0 else 0
        max_jornadas = total_estudiantes * 10
        pct_asistencia = round((jornadas_marcadas_total / max_jornadas) * 100, 2) if max_jornadas > 0 else 0

        resumen.append({
            'municipio': municipio,
            'total_estudiantes': total_estudiantes,
            'intentos_totales': intentos_totales,
            'pct_finalizacion_modulos': pct_completado,
            'pct_asistencia': pct_asistencia,
        })

    return resumen

@admin_bp.route('/dashboard/metrics', methods=['GET'])
@token_required
@admin_required
def get_dashboard_metrics(current_user):
    """Obtener mÃ©tricas generales del dashboard"""
    try:
        # Total de usuarios registrados
        total_usuarios = User.query.count()
        
        # Usuarios por estado
        usuarios_activos = User.query.filter_by(estado_cuenta='activa').count()
        usuarios_inactivos = User.query.filter_by(estado_cuenta='inactiva').count()
        usuarios_suspendidos = User.query.filter_by(estado_cuenta='suspendida').count()
        
        # Usuarios por rol
        estudiantes = User.query.filter_by(rol='estudiante').count()
        instructores = User.query.filter_by(rol='instructor').count()
        admins = User.query.filter_by(rol='admin').count()
        
        # Cursos activos
        cursos_activos = Curso.query.filter_by(activo=True).count()
        total_cursos = Curso.query.count()
        
        # Inscripciones (simulado por ahora)
        total_inscripciones = 0
        inscripciones_activas = 0
        
        # Usuarios registrados en los Ãºltimos 30 dÃ­as
        fecha_limite = datetime.utcnow() - timedelta(days=30)
        nuevos_usuarios = User.query.filter(User.fecha_creacion >= fecha_limite).count()
        
        municipios_resumen = []
        try:
            from src.routes.instructor import get_progreso_estudiantes
            progreso_payload = _extract_json_payload(get_progreso_estudiantes.__wrapped__.__wrapped__(current_user))
            progreso_estudiantes = progreso_payload.get('progreso_estudiantes') or []
            municipios_resumen = _build_municipios_summary(progreso_estudiantes)
        except Exception as resumen_error:
            print(f"Error generando resumen por municipio para admin: {resumen_error}")
        
        return jsonify({
            'usuarios': {
                'total': total_usuarios,
                'activos': usuarios_activos,
                'inactivos': usuarios_inactivos,
                'suspendidos': usuarios_suspendidos,
                'nuevos_30_dias': nuevos_usuarios
            },
            'roles': {
                'estudiantes': estudiantes,
                'instructores': instructores,
                'administradores': admins
            },
            'cursos': {
                'total': total_cursos,
                'activos': cursos_activos
            },
            'inscripciones': {
                'total': total_inscripciones,
                'activas': inscripciones_activas
            },
            'municipios_resumen': municipios_resumen
        }), 200
        
    except Exception as e:
        return jsonify({'error': f'Error al obtener mÃ©tricas: {str(e)}'}), 500


@admin_bp.route('/users/reset-progress', methods=['POST'])
@token_required
@admin_required
def reset_user_progress(current_user):
    """Resetear progreso de un estudiante a cero (para usuarios de prueba).
    Elimina: LogActividad (Completó/leccion_completada), IntentosEvaluacion,
    PuntosPlanNegocio, RespuestasPlanNegocio, AsistenciaJornada."""
    try:
        data = request.get_json() or {}
        email = (data.get('email') or '').strip()
        if not email:
            return jsonify({'success': False, 'error': 'Se requiere email en el body'}), 400

        user = User.query.filter_by(email=email).first()
        if not user:
            return jsonify({'success': False, 'error': f'Usuario con email {email} no encontrado'}), 404

        if user.rol not in ('estudiante', 'usuario'):
            return jsonify({'success': False, 'error': 'Solo se puede resetear progreso de estudiantes'}), 400

        user_id = user.id
        deleted = {
            'log_actividad': 0,
            'intentos_evaluacion': 0,
            'puntos_plan_negocio': 0,
            'respuestas_plan_negocio': 0,
            'asistencia_jornada': 0,
        }

        deleted['log_actividad'] = LogActividad.query.filter(
            LogActividad.usuario_id == user_id,
            or_(
                LogActividad.accion.like('Completó:%'),
                LogActividad.accion == 'leccion_completada'
            )
        ).delete(synchronize_session=False)

        deleted['intentos_evaluacion'] = IntentosEvaluacion.query.filter_by(usuario_id=user_id).delete(synchronize_session=False)
        deleted['puntos_plan_negocio'] = PuntosPlanNegocio.query.filter_by(usuario_id=user_id).delete(synchronize_session=False)
        deleted['respuestas_plan_negocio'] = RespuestasPlanNegocio.query.filter_by(usuario_id=user_id).delete(synchronize_session=False)
        deleted['asistencia_jornada'] = AsistenciaJornada.query.filter_by(estudiante_id=user_id).delete(synchronize_session=False)

        db.session.commit()

        return jsonify({
            'success': True,
            'message': f'Progreso de {email} reseteado a cero',
            'deleted': deleted
        }), 200

    except Exception as e:
        db.session.rollback()
        return jsonify({'success': False, 'error': str(e)}), 500


@admin_bp.route('/cupos/config', methods=['GET', 'PUT'])
@token_required
@admin_required
def cupos_config(current_user):
    try:
        if request.method == 'GET':
            cfg = CuposConfig.query.order_by(CuposConfig.id.desc()).first()
            return jsonify(cfg.to_dict() if cfg else {'modo': 'abierto', 'cupo_global_max': None, 'convocatoria': '2025'}), 200

        data = request.json or {}
        modo = (data.get('modo') or 'abierto').strip()
        if modo not in ['abierto', 'bloqueado']:
            return jsonify({'error': 'Modo invÃ¡lido'}), 400
        cupo_global_max = data.get('cupo_global_max')
        if cupo_global_max is not None:
            try:
                cupo_global_max = int(cupo_global_max)
                if cupo_global_max < 0:
                    return jsonify({'error': 'cupo_global_max debe ser >= 0'}), 400
            except Exception:
                return jsonify({'error': 'cupo_global_max debe ser entero'}), 400

        convocatoria = (data.get('convocatoria') or '2025').strip()

        cfg = CuposConfig(modo=modo, cupo_global_max=cupo_global_max, convocatoria=convocatoria)
        db.session.add(cfg)
        db.session.commit()
        try:
            db.session.add(LogActividad(
                usuario_id=current_user.id,
                accion='config_cupo_actualizada',
                detalles=f"modo={modo}, cupo_global_max={cupo_global_max}, convocatoria={convocatoria}"
            ))
            db.session.commit()
        except Exception:
            db.session.rollback()
        return jsonify({'message': 'ConfiguraciÃ³n actualizada', 'config': cfg.to_dict()}), 200
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': f'Error al actualizar configuraciÃ³n: {str(e)}'}), 500


@admin_bp.route('/cupos/municipios', methods=['GET', 'PUT'])
@token_required
@admin_required
def cupos_municipios(current_user):
    try:
        if request.method == 'GET':
            filas = MunicipioCupo.query.order_by(MunicipioCupo.subregion.asc(), MunicipioCupo.municipio_slug.asc()).all()
            return jsonify([f.to_dict() for f in filas]), 200

        data = request.json or {}
        items = data.get('items') or []
        with db.session.begin_nested():
            for it in items:
                slug = it.get('municipio_slug')
                cupo = it.get('cupo_max')
                if not slug:
                    continue
                try:
                    cupo = int(cupo)
                    if cupo < 0:
                        continue
                except Exception:
                    continue
                row = MunicipioCupo.query.filter_by(municipio_slug=slug).first()
                if row:
                    row.cupo_max = cupo
                else:
                    db.session.add(MunicipioCupo(municipio_slug=slug, subregion=it.get('subregion') or '', cupo_max=cupo))
        db.session.commit()
        try:
            db.session.add(LogActividad(
                usuario_id=current_user.id,
                accion='municipio_cupo_actualizado',
                detalles=f"items={len(items)}"
            ))
            db.session.commit()
        except Exception:
            db.session.rollback()
        return jsonify({'message': 'Cupos de municipios actualizados'}), 200
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': f'Error al actualizar cupos: {str(e)}'}), 500


@admin_bp.route('/cupos/estado', methods=['GET'])
@token_required
@admin_required
def cupos_estado(current_user):
    """Estado de cupos: global, por subregiÃ³n y por municipio"""
    try:
        # ConfiguraciÃ³n vigente
        cfg = CuposConfig.query.order_by(CuposConfig.id.desc()).first()
        convocatoria = cfg.convocatoria if cfg else '2025'

        # Cupos por municipio
        municipios = MunicipioCupo.query.all()
        muni_info = {m.municipio_slug: {'subregion': m.subregion, 'cupo_max': int(m.cupo_max)} for m in municipios}

        # Conteos por municipio
        q_base = db.session.query(User.municipio, User.estado_cuenta, db.func.count(User.id)).filter(
            User.convocatoria == convocatoria
        ).group_by(User.municipio, User.estado_cuenta).all()

        confirmados_map = {}
        espera_map = {}
        for muni, estado, cnt in q_base:
            if muni is None:
                continue
            if estado in ('activa', 'inactiva'):
                confirmados_map[muni] = confirmados_map.get(muni, 0) + int(cnt)
            elif estado == 'lista_espera':
                espera_map[muni] = espera_map.get(muni, 0) + int(cnt)

        # Construir lista municipal
        municipios_estado = []
        subregion_agg = {}
        total_cupo_municipal = 0
        total_confirmados = 0
        total_espera = 0
        for muni, info in muni_info.items():
            cupo_max = info['cupo_max']
            confirmados = confirmados_map.get(muni, 0)
            espera = espera_map.get(muni, 0)
            disponibles = max(0, cupo_max - confirmados)
            pct = (confirmados / cupo_max * 100.0) if cupo_max > 0 else 0.0
            municipios_estado.append({
                'municipio': muni,
                'subregion': info['subregion'],
                'cupo_max': cupo_max,
                'confirmados': confirmados,
                'lista_espera': espera,
                'disponibles': disponibles,
                'porcentaje': round(pct, 2),
            })

            sr = info['subregion']
            if sr not in subregion_agg:
                subregion_agg[sr] = {'cupo_max': 0, 'confirmados': 0, 'lista_espera': 0}
            subregion_agg[sr]['cupo_max'] += cupo_max
            subregion_agg[sr]['confirmados'] += confirmados
            subregion_agg[sr]['lista_espera'] += espera

            total_cupo_municipal += cupo_max
            total_confirmados += confirmados
            total_espera += espera

        subregiones_estado = []
        for sr, agg in subregion_agg.items():
            disp = max(0, agg['cupo_max'] - agg['confirmados'])
            pct = (agg['confirmados'] / agg['cupo_max'] * 100.0) if agg['cupo_max'] > 0 else 0.0
            subregiones_estado.append({
                'subregion': sr,
                'cupo_max': agg['cupo_max'],
                'confirmados': agg['confirmados'],
                'lista_espera': agg['lista_espera'],
                'disponibles': disp,
                'porcentaje': round(pct, 2),
            })

        # Global
        global_cupo = cfg.cupo_global_max if (cfg and cfg.cupo_global_max is not None) else total_cupo_municipal
        global_disponibles = max(0, global_cupo - total_confirmados)
        global_pct = (total_confirmados / global_cupo * 100.0) if global_cupo > 0 else 0.0

        return jsonify({
            'convocatoria': convocatoria,
            'global': {
                'cupo_max': int(global_cupo),
                'confirmados': int(total_confirmados),
                'lista_espera': int(total_espera),
                'disponibles': int(global_disponibles),
                'porcentaje': round(global_pct, 2),
            },
            'subregiones': subregiones_estado,
            'municipios': sorted(municipios_estado, key=lambda x: (x['subregion'], x['municipio']))
        }), 200
    except Exception as e:
        return jsonify({'error': f'Error al obtener estado de cupos: {str(e)}'}), 500


@admin_bp.route('/cupos/estado/export', methods=['GET'])
@token_required
@admin_required
def cupos_estado_export(current_user):
    """Exportar estado de cupos en Excel"""
    try:
        # Reusar lÃ³gica del endpoint anterior
        cfg = CuposConfig.query.order_by(CuposConfig.id.desc()).first()
        convocatoria = cfg.convocatoria if cfg else '2025'
        municipios = MunicipioCupo.query.all()
        muni_info = {m.municipio_slug: {'subregion': m.subregion, 'cupo_max': int(m.cupo_max)} for m in municipios}
        q_base = db.session.query(User.municipio, User.estado_cuenta, db.func.count(User.id)).filter(
            User.convocatoria == convocatoria
        ).group_by(User.municipio, User.estado_cuenta).all()
        confirmados_map, espera_map = {}, {}
        for muni, estado, cnt in q_base:
            if muni is None:
                continue
            if estado in ('activa', 'inactiva'):
                confirmados_map[muni] = confirmados_map.get(muni, 0) + int(cnt)
            elif estado == 'lista_espera':
                espera_map[muni] = espera_map.get(muni, 0) + int(cnt)
        
        # Crear workbook de Excel
        wb = Workbook()
        ws = wb.active
        ws.title = "Estado de Cupos"
        
        # Estilos para el encabezado
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        
        # Estilos para datos
        data_alignment = Alignment(horizontal="center", vertical="center")
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # Verde para disponible
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")    # Rojo para agotado
        yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid") # Amarillo para ocupaciÃ³n alta
        
        # Encabezados
        headers = ['SubregiÃ³n', 'Municipio', 'Cupo MÃ¡ximo', 'Confirmados', 'Lista Espera', 'Disponibles', '% OcupaciÃ³n']
        
        # Escribir encabezados con formato
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # Datos con formato condicional
        for row, (muni, info) in enumerate(sorted(muni_info.items(), key=lambda x: (x[1]['subregion'], x[0])), 2):
            cupo_max = info['cupo_max']
            conf = confirmados_map.get(muni, 0)
            esp = espera_map.get(muni, 0)
            disp = max(0, cupo_max - conf)
            pct = round((conf / cupo_max * 100.0), 2) if cupo_max > 0 else 0.0
            
            data = [info['subregion'], muni, cupo_max, conf, esp, disp, f"{pct}%"]
            
            # Escribir datos con formato condicional
            for col, value in enumerate(data, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.alignment = data_alignment
                cell.border = thin_border
                
                # Aplicar colores segÃºn ocupaciÃ³n (columna % OcupaciÃ³n)
                if col == 7:  # Columna de porcentaje
                    if pct >= 100:
                        cell.fill = red_fill  # Rojo si estÃ¡ al 100% o mÃ¡s
                    elif pct >= 80:
                        cell.fill = yellow_fill  # Amarillo si estÃ¡ entre 80-99%
                    else:
                        cell.fill = green_fill  # Verde si estÃ¡ por debajo del 80%
        
        # Ajustar ancho de columnas
        for col in range(1, len(headers) + 1):
            column_letter = get_column_letter(col)
            ws.column_dimensions[column_letter].width = 15
        
        # Crear archivo en memoria
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output.getvalue(), 200, {
            'Content-Type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'Content-Disposition': 'attachment; filename=estado_cupos.xlsx'
        }
    except Exception as e:
        return jsonify({'error': f'Error al exportar: {str(e)}'}), 500

@admin_bp.route('/users', methods=['GET'])
@token_required
@admin_or_evaluador_required
def get_users(current_user):
    """Obtener lista de usuarios con filtros"""
    try:
        page = request.args.get('page', 1, type=int)
        per_page = request.args.get('per_page', 20, type=int)
        estado = request.args.get('estado')
        rol = request.args.get('rol')
        search = request.args.get('search')
        estado_control = request.args.get('estado_control')
        municipio = request.args.get('municipio')
        fecha_inicio = request.args.get('fecha_inicio')
        fecha_fin = request.args.get('fecha_fin')
        solo_finalizados = request.args.get('solo_finalizados')
        
        query = User.query
        
        # Aplicar filtros
        if estado:
            query = query.filter_by(estado_cuenta=estado)
        if rol:
            query = query.filter_by(rol=rol)
        if estado_control:
            query = query.filter_by(estado_control=estado_control)
        if municipio:
            query = query.filter_by(municipio=municipio)
        if search:
            search_term = f"%{search}%"
            query = query.filter(
                (User.nombre.ilike(search_term)) |
                (User.apellido.ilike(search_term)) |
                (User.email.ilike(search_term)) |
                (User.numero_documento.ilike(search_term))
            )
        
        # Filtros de fecha
        if fecha_inicio:
            try:
                fecha_inicio_dt = datetime.fromisoformat(fecha_inicio.replace('T', ' '))
                query = query.filter(User.fecha_creacion >= fecha_inicio_dt)
            except ValueError:
                pass  # Ignorar fechas invÃ¡lidas
        
        if fecha_fin:
            try:
                fecha_fin_dt = datetime.fromisoformat(fecha_fin.replace('T', ' '))
                query = query.filter(User.fecha_creacion <= fecha_fin_dt)
            except ValueError:
                pass  # Ignorar fechas invÃ¡lidas
        
        # Filtro de finalizaciÃ³n
        if solo_finalizados == 'finalizados':
            query = query.filter(User.fecha_finalizacion.isnot(None))
        elif solo_finalizados == 'no_finalizados':
            query = query.filter(User.fecha_finalizacion.is_(None))
        
        # Ordenar por fecha de creaciÃ³n (mÃ¡s recientes primero)
        query = query.order_by(User.fecha_creacion.desc())
        
        # PaginaciÃ³n
        pagination = query.paginate(
            page=page, 
            per_page=per_page, 
            error_out=False
        )
        
        users = pagination.items
        
        return jsonify({
            'users': [user.to_dict() for user in users],
            'pagination': {
                'page': page,
                'per_page': per_page,
                'total': pagination.total,
                'pages': pagination.pages,
                'has_next': pagination.has_next,
                'has_prev': pagination.has_prev
            }
        }), 200
        
    except Exception as e:
        return jsonify({'error': f'Error al obtener usuarios: {str(e)}'}), 500

@admin_bp.route('/users/<int:user_id>', methods=['PUT'])
@token_required
@admin_or_evaluador_required
def update_user(current_user, user_id):
    """Actualizar usuario (estado, rol, etc.)"""
    try:
        user = User.query.get_or_404(user_id)
        data = request.json
        
        # Campos permitidos para actualizaciÃ³n
        if 'estado_cuenta' in data:
            user.estado_cuenta = data['estado_cuenta']
        if 'rol' in data:
            user.rol = data['rol']
        if 'nombre' in data:
            user.nombre = data['nombre']
        if 'apellido' in data:
            user.apellido = data['apellido']
        
        user.fecha_actualizacion = datetime.utcnow()
        
        # Registrar actividad (temporalmente comentado por problemas de mapeo)
        # log = LogActividad(
        #     usuario_id=session['user_id'],
        #     accion='actualizar_usuario',
        #     detalles=f'Usuario {user.email} actualizado por administrador'
        # )
        # db.session.add(log)
        
        db.session.commit()
        
        return jsonify({
            'message': 'Usuario actualizado exitosamente',
            'user': user.to_dict()
        }), 200
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': f'Error al actualizar usuario: {str(e)}'}), 500

@admin_bp.route('/users/bulk-update', methods=['POST'])
@token_required
@admin_or_evaluador_required
def bulk_update_users(current_user):
    """ActualizaciÃ³n masiva de usuarios"""
    try:
        data = request.json
        user_ids = data.get('user_ids', [])
        updates = data.get('updates', {})
        
        if not user_ids:
            return jsonify({'error': 'No se proporcionaron IDs de usuarios'}), 400
        
        users = User.query.filter(User.id.in_(user_ids)).all()
        
        for user in users:
            if 'estado_cuenta' in updates:
                user.estado_cuenta = updates['estado_cuenta']
            if 'rol' in updates:
                user.rol = updates['rol']
            
            user.fecha_actualizacion = datetime.utcnow()
        
        # Registrar actividad (temporalmente comentado por problemas de mapeo)
        # log = LogActividad(
        #     usuario_id=session['user_id'],
        #     accion='actualizacion_masiva_usuarios',
        #     detalles=f'ActualizaciÃ³n masiva de {len(users)} usuarios'
        # )
        # db.session.add(log)
        
        db.session.commit()
        
        return jsonify({
            'message': f'{len(users)} usuarios actualizados exitosamente'
        }), 200
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': f'Error en actualizaciÃ³n masiva: {str(e)}'}), 500

@admin_bp.route('/users/import', methods=['POST'])
@token_required
@admin_required
def import_users(current_user):
    """Importar usuarios desde CSV"""
    try:
        if 'file' not in request.files:
            return jsonify({'error': 'No se proporcionÃ³ archivo'}), 400
        
        file = request.files['file']
        if file.filename == '':
            return jsonify({'error': 'No se seleccionÃ³ archivo'}), 400
        
        if not file.filename.endswith('.csv'):
            return jsonify({'error': 'El archivo debe ser CSV'}), 400
        
        # Leer CSV
        stream = io.StringIO(file.stream.read().decode("UTF8"), newline=None)
        csv_reader = csv.DictReader(stream)
        
        imported_count = 0
        errors = []
        
        for row_num, row in enumerate(csv_reader, 2):  # Empezar en 2 porque la fila 1 es header
            try:
                # Validar campos requeridos
                required_fields = ['nombre', 'apellido', 'email', 'password', 'tipo_documento', 'numero_documento']
                for field in required_fields:
                    if not row.get(field):
                        errors.append(f"Fila {row_num}: Campo '{field}' es requerido")
                        continue
                
                # Verificar si el email ya existe
                if User.query.filter_by(email=row['email']).first():
                    errors.append(f"Fila {row_num}: Email '{row['email']}' ya existe")
                    continue
                
                # Crear usuario
                user = User(
                    nombre=row['nombre'],
                    apellido=row['apellido'],
                    email=row['email'],
                    tipo_documento=row['tipo_documento'],
                    numero_documento=row['numero_documento'],
                    rol=row.get('rol', 'estudiante'),
                    estado_cuenta=row.get('estado_cuenta', 'inactiva')
                )
                user.set_password(row['password'])
                
                db.session.add(user)
                imported_count += 1
                
            except Exception as e:
                errors.append(f"Fila {row_num}: Error - {str(e)}")
        
        if imported_count > 0:
            # Registrar actividad (temporalmente comentado por problemas de mapeo)
            # log = LogActividad(
            #     usuario_id=session['user_id'],
            #     accion='importar_usuarios',
            #     detalles=f'Importados {imported_count} usuarios desde CSV'
            # )
            # db.session.add(log)
            
            db.session.commit()
        
        return jsonify({
            'message': f'ImportaciÃ³n completada',
            'imported_count': imported_count,
            'errors': errors
        }), 200
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': f'Error en importaciÃ³n: {str(e)}'}), 500

@admin_bp.route('/users/download-excel', methods=['GET'])
@token_required
@admin_or_evaluador_required
def download_users_excel(current_user):
    """Exportar usuarios a Excel - Mismo formato que muestra 'Ver Detalles'"""
    try:
        # Obtener todos los usuarios (solo rol 'usuario', no admin ni instructor)
        users = User.query.filter_by(rol='usuario').all()
        
        # Crear workbook de Excel
        wb = Workbook()
        ws = wb.active
        ws.title = "Usuarios"
        
        # Estilos
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        data_alignment = Alignment(horizontal="left", vertical="center")
        
        # Encabezados - Exactamente como aparecen en "Ver Detalles"
        headers = [
            'Nombre', 'Email', 'Documento', 'Fecha Nacimiento', 'Sexo', 
            'Estado Civil', 'TelÃ©fono', 'DirecciÃ³n', 'Municipio', 
            'Corregimiento/Vereda', 'Tipo Persona', 'Convocatoria',
            'Estado', 'Fecha Registro', 'Fecha FinalizaciÃ³n'
        ]
        
        # Escribir encabezados
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # Configurar zona horaria Colombia
        from pytz import timezone as pytz_timezone
        colombia_tz = pytz_timezone('America/Bogota')
        utc_tz = pytz_timezone('UTC')
        
        # Escribir datos de usuarios
        for row_idx, user in enumerate(users, 2):
            # Formatear fechas en hora de Colombia (UTC-5)
            fecha_registro = ''
            if user.fecha_creacion:
                try:
                    if user.fecha_creacion.tzinfo is None:
                        fecha_utc = utc_tz.localize(user.fecha_creacion)
                    else:
                        fecha_utc = user.fecha_creacion
                    fecha_col = fecha_utc.astimezone(colombia_tz)
                    fecha_registro = fecha_col.strftime('%Y-%m-%d %H:%M:%S')
                except:
                    fecha_registro = str(user.fecha_creacion)
            
            fecha_finalizacion = ''
            if user.fecha_finalizacion:
                try:
                    if user.fecha_finalizacion.tzinfo is None:
                        fecha_fin_utc = utc_tz.localize(user.fecha_finalizacion)
                    else:
                        fecha_fin_utc = user.fecha_finalizacion
                    fecha_fin_col = fecha_fin_utc.astimezone(colombia_tz)
                    fecha_finalizacion = fecha_fin_col.strftime('%Y-%m-%d %H:%M:%S')
                except:
                    fecha_finalizacion = str(user.fecha_finalizacion)
            
            # Datos en el mismo orden que los encabezados
            row_data = [
                f"{user.nombre or ''} {user.apellido or ''}".strip(),
                user.email or '',
                f"{user.tipo_documento or ''} - {user.numero_documento or ''}".strip(' -'),
                user.fecha_nacimiento or '',
                user.sexo or '',
                user.estado_civil or '',
                user.telefono or '',
                user.direccion or '',
                user.municipio or '',
                user.corregimiento_vereda or 'No especificado',
                user.tipo_persona or '',
                user.convocatoria or '',
                user.estado_cuenta or '',
                fecha_registro,
                fecha_finalizacion
            ]
            
            # Escribir fila
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.alignment = data_alignment
                cell.border = thin_border
        
        # Ajustar ancho de columnas
        ws.column_dimensions['A'].width = 30  # Nombre
        ws.column_dimensions['B'].width = 30  # Email
        ws.column_dimensions['C'].width = 20  # Documento
        ws.column_dimensions['D'].width = 15  # Fecha Nacimiento
        ws.column_dimensions['E'].width = 12  # Sexo
        ws.column_dimensions['F'].width = 15  # Estado Civil
        ws.column_dimensions['G'].width = 15  # TelÃ©fono
        ws.column_dimensions['H'].width = 30  # DirecciÃ³n
        ws.column_dimensions['I'].width = 20  # Municipio
        ws.column_dimensions['J'].width = 25  # Corregimiento/Vereda
        ws.column_dimensions['K'].width = 15  # Tipo Persona
        ws.column_dimensions['L'].width = 15  # Convocatoria
        ws.column_dimensions['M'].width = 15  # Estado
        ws.column_dimensions['N'].width = 20  # Fecha Registro
        ws.column_dimensions['O'].width = 20  # Fecha FinalizaciÃ³n
        
        # Guardar en memoria
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        # Log de actividad
        try:
            log = LogActividad(
                usuario_id=current_user.id,
                accion='descargar_usuarios_excel',
                detalles=f'Descargados {len(users)} usuarios en Excel'
            )
            db.session.add(log)
            db.session.commit()
        except:
            pass
        
        # Retornar archivo
        from flask import Response
        return Response(
            output.getvalue(),
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': 'attachment; filename=usuarios.xlsx'}
        )
        
    except Exception as e:
        return jsonify({'error': f'Error en exportaciÃ³n: {str(e)}'}), 500

@admin_bp.route('/users/<int:user_id>/certificados', methods=['PUT'])
@token_required
@admin_or_evaluador_required
def actualizar_resultado_certificados(current_user, user_id):
    """Actualizar el resultado de los certificados de control de un usuario"""
    try:
        data = request.json
        resultado = data.get('resultado_certificados')
        
        # Validar resultado
        if resultado not in ['pendiente', 'limpio', 'inhabilidad_detectada']:
            return jsonify({'error': 'Resultado invÃ¡lido. Debe ser: pendiente, limpio o inhabilidad_detectada'}), 400
        
        user = User.query.get_or_404(user_id)
        resultado_anterior = user.resultado_certificados
        
        # Actualizar resultado
        user.resultado_certificados = resultado
        
        # REGLA AUTOMÃTICA DE RECHAZO
        if resultado == 'inhabilidad_detectada':
            estado_anterior = user.estado_cuenta
            user.estado_cuenta = 'rechazada'
            
            # Log de auditorÃ­a del rechazo automÃ¡tico
            detalle_rechazo = f"RECHAZO AUTOMÃTICO: Usuario {user.nombre} {user.apellido} (ID: {user.id}) rechazado automÃ¡ticamente por inhabilidad detectada en certificados de control. Estado anterior: {estado_anterior}. Resultado anterior: {resultado_anterior}"
            
            db.session.add(LogActividad(
                usuario_id=session['user_id'],
                accion='rechazo_automatico_certificados',
                detalles=detalle_rechazo
            ))
            
            # Log tambiÃ©n en el usuario afectado
            db.session.add(LogActividad(
                usuario_id=user.id,
                accion='cuenta_rechazada_automaticamente',
                detalles=f"Cuenta rechazada automÃ¡ticamente por inhabilidad detectada en certificados de control. Administrador que marcÃ³: {session.get('user_email', 'N/A')}"
            ))
            
        elif resultado == 'limpio' and user.estado_cuenta == 'rechazada':
            # Si se marca como limpio y estaba rechazada, revertir a inactiva para revisiÃ³n manual
            user.estado_cuenta = 'inactiva'
            
            db.session.add(LogActividad(
                usuario_id=session['user_id'],
                accion='revertir_rechazo_certificados',
                detalles=f"Rechazo revertido para usuario {user.nombre} {user.apellido} (ID: {user.id}). Certificados marcados como limpios. Estado cambiado a inactiva para revisiÃ³n."
            ))
        
        db.session.commit()
        
        return jsonify({
            'message': f'Resultado de certificados actualizado a: {resultado}',
            'usuario': user.to_dict(),
            'rechazo_automatico_aplicado': resultado == 'inhabilidad_detectada'
        }), 200
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': f'Error al actualizar certificados: {str(e)}'}), 500

@admin_bp.route('/certificados/estadisticas', methods=['GET'])
@token_required
@admin_or_evaluador_required
def estadisticas_certificados(current_user):
    """Obtener estadÃ­sticas de certificados de control y rechazos automÃ¡ticos"""
    try:
        # EstadÃ­sticas generales
        total_usuarios = User.query.count()
        control_completo = User.query.filter_by(estado_control='completo').count()
        control_pendiente = User.query.filter_by(estado_control='pendiente').count()
        
        # Resultados de certificados
        resultado_pendiente = User.query.filter_by(resultado_certificados='pendiente').count()
        resultado_limpio = User.query.filter_by(resultado_certificados='limpio').count()
        resultado_inhabilidad = User.query.filter_by(resultado_certificados='inhabilidad_detectada').count()
        
        # Rechazos automÃ¡ticos
        rechazados_automaticamente = User.query.filter_by(estado_cuenta='rechazada', resultado_certificados='inhabilidad_detectada').count()
        
        # Logs de rechazos automÃ¡ticos recientes
        rechazos_recientes = LogActividad.query.filter_by(accion='rechazo_automatico_certificados').order_by(LogActividad.fecha.desc()).limit(10).all()
        
        return jsonify({
            'estadisticas_generales': {
                'total_usuarios': total_usuarios,
                'control_completo': control_completo,
                'control_pendiente': control_pendiente
            },
            'resultados_certificados': {
                'pendiente': resultado_pendiente,
                'limpio': resultado_limpio,
                'inhabilidad_detectada': resultado_inhabilidad
            },
            'rechazos_automaticos': {
                'total': rechazados_automaticamente,
                'porcentaje': round((rechazados_automaticamente / total_usuarios * 100), 2) if total_usuarios > 0 else 0
            },
            'rechazos_recientes': [log.to_dict() for log in rechazos_recientes]
        }), 200
        
    except Exception as e:
        return jsonify({'error': f'Error al obtener estadÃ­sticas: {str(e)}'}), 500

@admin_bp.route('/logs', methods=['GET'])
@token_required
@admin_required
def get_logs(current_user):
    """Obtener logs de actividad"""
    try:
        page = request.args.get('page', 1, type=int)
        per_page = request.args.get('per_page', 50, type=int)
        usuario_id = request.args.get('usuario_id', type=int)
        accion = request.args.get('accion')
        
        # Por ahora retornamos logs vacÃ­os para evitar errores
        logs = []
        
        return jsonify({
            'logs': logs,
            'pagination': {
                'page': page,
                'per_page': per_page,
                'total': 0,
                'pages': 1,
                'has_next': False,
                'has_prev': False
            }
        }), 200
        
    except Exception as e:
        return jsonify({'error': f'Error al obtener logs: {str(e)}'}), 500 

@admin_bp.route('/courses', methods=['GET'])
@token_required
@admin_required
def get_courses(current_user):
    """Endpoint temporal - devuelve lista vacÃ­a para evitar errores"""
    return jsonify({
        'courses': [],
        'pagination': {
            'page': 1,
            'pages': 1,
            'per_page': 20,
            'total': 0,
            'has_next': False,
            'has_prev': False
        }
    }), 200

@admin_bp.route('/courses', methods=['POST'])
@token_required
@admin_required
def create_course(current_user):
    """Crear nuevo curso"""
    try:
        data = request.json
        
        # Validar campos obligatorios
        if not data.get('titulo'):
            return jsonify({'error': 'El tÃ­tulo del curso es obligatorio'}), 400
        
        # Procesar instructor_id
        instructor_id = data.get('instructor_id')
        if instructor_id == '' or instructor_id is None:
            instructor_id = None
        else:
            try:
                instructor_id = int(instructor_id)
            except (ValueError, TypeError):
                return jsonify({'error': 'ID de instructor invÃ¡lido'}), 400
        
        # Crear nuevo curso
        curso = Curso(
            titulo=data['titulo'],
            descripcion=data.get('descripcion', ''),
            instructor_id=instructor_id,
            estado=data.get('estado', 'activo'),
            fecha_apertura=datetime.fromisoformat(data['fecha_apertura']) if data.get('fecha_apertura') else None,
            fecha_cierre=datetime.fromisoformat(data['fecha_cierre']) if data.get('fecha_cierre') else None,
            duracion_horas=data.get('duracion_horas', 0),
            nivel=data.get('nivel', 'bÃ¡sico'),
            categoria=data.get('categoria'),
            imagen_url=data.get('imagen_url'),
            max_estudiantes=data.get('max_estudiantes', 0)
        )
        
        db.session.add(curso)
        db.session.commit()
        
        # Obtener informaciÃ³n del curso creado con instructor
        course_dict = curso.to_dict()
        if curso.instructor_id:
            instructor = User.query.get(curso.instructor_id)
            if instructor:
                course_dict['instructor_nombre'] = f"{instructor.nombre} {instructor.apellido}"
            else:
                course_dict['instructor_nombre'] = 'Instructor no encontrado'
        else:
            course_dict['instructor_nombre'] = 'Sin asignar'
        
        return jsonify({
            'message': 'Curso creado exitosamente',
            'course': course_dict
        }), 201
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': f'Error al crear curso: {str(e)}'}), 500

@admin_bp.route('/courses/<int:course_id>', methods=['GET'])
@token_required
@admin_required
def get_course(current_user, course_id):
    """Obtener curso especÃ­fico"""
    try:
        curso = Curso.query.get_or_404(course_id)
        course_dict = curso.to_dict()
        
        # Obtener informaciÃ³n del instructor
        if curso.instructor_id:
            instructor = User.query.get(curso.instructor_id)
            if instructor:
                course_dict['instructor_nombre'] = f"{instructor.nombre} {instructor.apellido}"
            else:
                course_dict['instructor_nombre'] = 'Instructor no encontrado'
        else:
            course_dict['instructor_nombre'] = 'Sin asignar'
        
        # Obtener total de inscripciones
        total_inscripciones = Inscripcion.query.filter_by(curso_id=curso.id).count()
        course_dict['total_inscripciones'] = total_inscripciones
        
        return jsonify(course_dict), 200
        
    except Exception as e:
        return jsonify({'error': f'Error al obtener curso: {str(e)}'}), 500

@admin_bp.route('/courses/<int:course_id>', methods=['PUT'])
@token_required
@admin_required
def update_course(current_user, course_id):
    """Actualizar curso"""
    try:
        curso = Curso.query.get_or_404(course_id)
        data = request.json
        
        # Actualizar campos
        if 'titulo' in data:
            curso.titulo = data['titulo']
        if 'descripcion' in data:
            curso.descripcion = data['descripcion']
        if 'instructor_id' in data:
            # Procesar instructor_id
            instructor_id = data['instructor_id']
            if instructor_id == '' or instructor_id is None:
                curso.instructor_id = None
            else:
                try:
                    curso.instructor_id = int(instructor_id)
                except (ValueError, TypeError):
                    return jsonify({'error': 'ID de instructor invÃ¡lido'}), 400
        if 'estado' in data:
            curso.estado = data['estado']
        if 'fecha_apertura' in data:
            curso.fecha_apertura = datetime.fromisoformat(data['fecha_apertura']) if data['fecha_apertura'] else None
        if 'fecha_cierre' in data:
            curso.fecha_cierre = datetime.fromisoformat(data['fecha_cierre']) if data['fecha_cierre'] else None
        if 'duracion_horas' in data:
            curso.duracion_horas = data['duracion_horas']
        if 'nivel' in data:
            curso.nivel = data['nivel']
        if 'categoria' in data:
            curso.categoria = data['categoria']
        if 'imagen_url' in data:
            curso.imagen_url = data['imagen_url']
        if 'max_estudiantes' in data:
            curso.max_estudiantes = data['max_estudiantes']
        
        curso.fecha_actualizacion = datetime.utcnow()
        
        db.session.commit()
        
        # Obtener informaciÃ³n actualizada del curso con instructor
        course_dict = curso.to_dict()
        if curso.instructor_id:
            instructor = User.query.get(curso.instructor_id)
            if instructor:
                course_dict['instructor_nombre'] = f"{instructor.nombre} {instructor.apellido}"
            else:
                course_dict['instructor_nombre'] = 'Instructor no encontrado'
        else:
            course_dict['instructor_nombre'] = 'Sin asignar'
        
        return jsonify({
            'message': 'Curso actualizado exitosamente',
            'course': course_dict
        }), 200
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': f'Error al actualizar curso: {str(e)}'}), 500

@admin_bp.route('/courses/<int:course_id>', methods=['DELETE'])
@token_required
@admin_required
def delete_course(current_user, course_id):
    """Eliminar curso"""
    try:
        curso = Curso.query.get_or_404(course_id)
        
        # Verificar si hay inscripciones
        if curso.inscripciones.count() > 0:
            return jsonify({'error': 'No se puede eliminar un curso que tiene estudiantes inscritos'}), 400
        
        db.session.delete(curso)
        db.session.commit()
        
        return jsonify({'message': 'Curso eliminado exitosamente'}), 200
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': f'Error al eliminar curso: {str(e)}'}), 500

@admin_bp.route('/instructors', methods=['GET'])
@token_required
@admin_required
def get_instructors(current_user):
    """Obtener lista de instructores"""
    try:
        instructores = User.query.filter_by(rol='instructor').all()
        return jsonify([{
            'id': user.id,
            'nombre': f"{user.nombre} {user.apellido}",
            'email': user.email,
            'estado_cuenta': user.estado_cuenta
        } for user in instructores]), 200
        
    except Exception as e:
        return jsonify({'error': f'Error al obtener instructores: {str(e)}'}), 500 

@admin_bp.route('/users/<int:user_id>/documento', methods=['GET'])
@token_required
@admin_required
def download_user_documento(current_user, user_id):
    """Descargar/ver documento de identidad PDF del usuario"""
    try:
        user = User.query.get_or_404(user_id)
        if not user.documento_pdf:
            return jsonify({'error': 'El usuario no tiene documento PDF cargado'}), 404
        
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.pdf')
        tmp.write(user.documento_pdf)
        tmp.flush()
        tmp.seek(0)
        
        filename = user.documento_pdf_nombre or 'documento.pdf'
        return send_file(tmp.name, mimetype='application/pdf', as_attachment=False, download_name=filename)
    except Exception as e:
        return jsonify({'error': f'Error al descargar documento: {str(e)}'}), 500

@admin_bp.route('/users/<int:user_id>/requisitos', methods=['GET'])
@token_required
@admin_required
def download_user_requisitos(current_user, user_id):
    """Descargar/ver requisitos PDF del usuario"""
    try:
        user = User.query.get_or_404(user_id)
        if not user.requisitos_pdf:
            return jsonify({'error': 'El usuario no tiene requisitos PDF cargados'}), 404
        
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.pdf')
        tmp.write(user.requisitos_pdf)
        tmp.flush()
        tmp.seek(0)
        
        filename = user.requisitos_pdf_nombre or 'requisitos.pdf'
        return send_file(tmp.name, mimetype='application/pdf', as_attachment=False, download_name=filename)
    except Exception as e:
        return jsonify({'error': f'Error al descargar requisitos: {str(e)}'}), 500

# ===== ENDPOINTS PARA GESTIÃ“N DE FASES =====

@admin_bp.route('/stats/phases', methods=['GET'])
@token_required
@admin_required
def get_phase_stats(current_user):
    """Obtener estadÃ­sticas de fases de usuarios"""
    try:
        # EstadÃ­sticas generales
        total_users = User.query.count()
        
        # Usuarios por fase
        inscripcion = User.query.filter_by(fase_actual='inscripcion').count()
        formacion = User.query.filter_by(fase_actual='formacion').count()
        entrega_activos = User.query.filter_by(fase_actual='entrega_activos').count()
        
        # Usuarios con fases completadas
        fases_completadas = User.query.filter_by(fase_completada=True).count()
        
        # Progreso por fase
        inscripcion_completada = User.query.filter_by(fase_actual='inscripcion', fase_completada=True).count()
        formacion_completada = User.query.filter_by(fase_actual='formacion', fase_completada=True).count()
        entrega_completada = User.query.filter_by(fase_actual='entrega_activos', fase_completada=True).count()
        
        return jsonify({
            'total_users': total_users,
            'inscripcion': inscripcion,
            'formacion': formacion,
            'entrega_activos': entrega_activos,
            'fases_completadas': fases_completadas,
            'progreso': {
                'inscripcion': {
                    'total': inscripcion,
                    'completada': inscripcion_completada,
                    'porcentaje': round((inscripcion_completada / inscripcion * 100) if inscripcion > 0 else 0, 2)
                },
                'formacion': {
                    'total': formacion,
                    'completada': formacion_completada,
                    'porcentaje': round((formacion_completada / formacion * 100) if formacion > 0 else 0, 2)
                },
                'entrega_activos': {
                    'total': entrega_activos,
                    'completada': entrega_completada,
                    'porcentaje': round((entrega_completada / entrega_activos * 100) if entrega_activos > 0 else 0, 2)
                }
            }
        }), 200
        
    except Exception as e:
        return jsonify({'error': f'Error al obtener estadÃ­sticas de fases: {str(e)}'}), 500

@admin_bp.route('/users', methods=['GET'])
@token_required
@admin_required
def get_all_users(current_user):
    """Obtener todos los usuarios con informaciÃ³n de fases"""
    try:
        users = User.query.all()
        
        users_data = []
        for user in users:
            users_data.append({
                'id': user.id,
                'nombre': user.nombre,
                'apellido': user.apellido,
                'email': user.email,
                'telefono': user.telefono,
                'municipio': user.municipio,
                'emprendimiento_nombre': user.emprendimiento_nombre,
                'emprendimiento_sector': user.emprendimiento_sector,
                'tipo_persona': user.tipo_persona,
                'fase_actual': getattr(user, 'fase_actual', 'inscripcion'),
                'fecha_entrada_fase': user.fecha_entrada_fase.isoformat() if hasattr(user, 'fecha_entrada_fase') and user.fecha_entrada_fase else None,
                'fase_completada': getattr(user, 'fase_completada', False),
                'estado_inscripcion': user.estado_inscripcion,
                'formulario_enviado': user.formulario_enviado,
                'fecha_creacion': user.fecha_creacion.isoformat() if user.fecha_creacion else None,
                'estado_cuenta': user.estado_cuenta,
                'rol': user.rol
            })
        
        return jsonify({'users': users_data}), 200
        
    except Exception as e:
        return jsonify({'error': f'Error al obtener usuarios: {str(e)}'}), 500

@admin_bp.route('/users/all', methods=['GET'])
@token_required
@admin_or_evaluador_required
def get_all_users_complete(current_user):
    """Obtener usuarios con paginaciÃ³n para evitar lÃ­mite de 6MB de Lambda"""
    try:
        # Obtener parÃ¡metros de paginaciÃ³n
        page = request.args.get('page', 1, type=int)
        per_page = request.args.get('per_page', 2000, type=int)  # 2000 usuarios por pÃ¡gina
        
        # Validar parÃ¡metros
        if page < 1:
            page = 1
        if per_page < 1 or per_page > 2000:
            per_page = 2000
        
        # Calcular offset
        offset = (page - 1) * per_page
        
        # Usar consulta SQL directa para mejor rendimiento
        # Solo seleccionar campos necesarios, evitando campos LargeBinary pesados
        query = db.session.query(
            User.id,
            User.nombre,
            User.apellido,
            User.email,
            User.telefono,
            User.fecha_nacimiento,
            User.sexo,
            User.estado_civil,
            User.direccion,
            User.municipio,
            User.corregimiento_vereda,
            User.pais,
            User.ciudad,
            User.tipo_documento,
            User.numero_documento,
            # Solo nombres de archivos, no los archivos binarios
            User.doc_terminos_pdf_nombre,
            User.doc_uso_imagen_pdf_nombre,
            User.doc_plan_negocio_nombre,
            User.doc_vecindad_pdf_nombre,
            User.video_url,
            User.rut_pdf_nombre,
            User.cedula_pdf_nombre,
            User.cedula_representante_pdf_nombre,
            User.cert_existencia_pdf_nombre,
            User.ruv_pdf_nombre,
            User.sisben_pdf_nombre,
            User.grupo_etnico_pdf_nombre,
            User.arn_pdf_nombre,
            User.discapacidad_pdf_nombre,
            User.antecedentes_fiscales_pdf_nombre,
            User.antecedentes_disciplinarios_pdf_nombre,
            User.antecedentes_judiciales_pdf_nombre,
            User.redam_pdf_nombre,
            User.inhabilidades_sexuales_pdf_nombre,
            User.declaracion_capacidad_legal_pdf_nombre,
            User.estado_control,
            User.resultado_certificados,
            User.emprendimiento_formalizado,
            User.matricula_mercantil_pdf_nombre,
            User.facturas_6meses_pdf_nombre,
            User.publicaciones_redes_pdf_nombre,
            User.registro_ventas_pdf_nombre,
            User.mujer_cabeza_familia,
            User.victima_conflicto,
            User.persona_discapacidad,
            User.pertenencia_etnica,
            User.sisben_grupo,
            User.persona_reincorporacion,
            User.tiempo_funcionamiento,
            User.empleos_generados,
            User.acceso_mercados,
            User.financiado_estado,
            User.financiado_regalias,
            User.financiado_camara_comercio,
            User.financiado_incubadoras,
            User.financiado_otro,
            User.financiado_otro_texto,
            User.declara_veraz,
            User.declara_no_beneficiario,
            User.acepta_terminos,
            User.fecha_aceptacion_terminos,
            User.estado_inscripcion,
            User.paso_actual,
            User.fecha_ultimo_guardado,
            User.formulario_enviado,
            User.rol,
            User.fecha_creacion,
            User.fecha_actualizacion,
            User.fecha_finalizacion,
            User.estado_cuenta,
            User.convocatoria,
            User.emprendimiento_nombre,
            User.emprendimiento_sector,
            User.tipo_persona
        ).order_by(User.fecha_creacion.desc())
        
        # Obtener total de usuarios para paginaciÃ³n
        total_users = query.count()
        
        # Aplicar paginaciÃ³n
        results = query.limit(per_page).offset(offset).all()
        
        # Calcular metadatos de paginaciÃ³n
        total_pages = (total_users + per_page - 1) // per_page
        has_next = page < total_pages
        has_prev = page > 1
        
        # Procesar resultados
        users_data = []
        for row in results:
            try:
                # Contar documentos de forma rÃ¡pida (solo campos _nombre que no sean None/vacÃ­os)
                doc_count = sum(1 for field in [
                    row.doc_terminos_pdf_nombre, row.doc_uso_imagen_pdf_nombre, 
                    row.doc_plan_negocio_nombre, row.doc_vecindad_pdf_nombre,
                    row.rut_pdf_nombre, row.cedula_pdf_nombre, 
                    row.cedula_representante_pdf_nombre, row.cert_existencia_pdf_nombre,
                    row.ruv_pdf_nombre, row.sisben_pdf_nombre, 
                    row.grupo_etnico_pdf_nombre, row.arn_pdf_nombre,
                    row.discapacidad_pdf_nombre, row.antecedentes_fiscales_pdf_nombre,
                    row.antecedentes_disciplinarios_pdf_nombre, row.antecedentes_judiciales_pdf_nombre,
                    row.redam_pdf_nombre, row.inhabilidades_sexuales_pdf_nombre,
                    row.declaracion_capacidad_legal_pdf_nombre, row.matricula_mercantil_pdf_nombre,
                    row.facturas_6meses_pdf_nombre, row.publicaciones_redes_pdf_nombre,
                    row.registro_ventas_pdf_nombre
                ] if field and field.strip())
                
                # Agregar video si existe
                if row.video_url and row.video_url.strip():
                    doc_count += 1
                
                user_data = {
                    'id': row.id,
                    'nombre': row.nombre,
                    'apellido': row.apellido,
                    'email': row.email,
                    'telefono': row.telefono,
                    'fecha_nacimiento': row.fecha_nacimiento.isoformat() if row.fecha_nacimiento else None,
                    'sexo': row.sexo,
                    'estado_civil': row.estado_civil,
                    'direccion': row.direccion,
                    'municipio': row.municipio,
                    'corregimiento_vereda': row.corregimiento_vereda,
                    'pais': row.pais,
                    'ciudad': row.ciudad,
                    'tipo_documento': row.tipo_documento,
                    'numero_documento': row.numero_documento,
                    'doc_terminos_pdf_nombre': row.doc_terminos_pdf_nombre,
                    'doc_uso_imagen_pdf_nombre': row.doc_uso_imagen_pdf_nombre,
                    'doc_plan_negocio_nombre': row.doc_plan_negocio_nombre,
                    'doc_vecindad_pdf_nombre': row.doc_vecindad_pdf_nombre,
                    'video_url': row.video_url,
                    'rut_pdf_nombre': row.rut_pdf_nombre,
                    'cedula_pdf_nombre': row.cedula_pdf_nombre,
                    'cedula_representante_pdf_nombre': row.cedula_representante_pdf_nombre,
                    'cert_existencia_pdf_nombre': row.cert_existencia_pdf_nombre,
                    'ruv_pdf_nombre': row.ruv_pdf_nombre,
                    'sisben_pdf_nombre': row.sisben_pdf_nombre,
                    'grupo_etnico_pdf_nombre': row.grupo_etnico_pdf_nombre,
                    'arn_pdf_nombre': row.arn_pdf_nombre,
                    'discapacidad_pdf_nombre': row.discapacidad_pdf_nombre,
                    'antecedentes_fiscales_pdf_nombre': row.antecedentes_fiscales_pdf_nombre,
                    'antecedentes_disciplinarios_pdf_nombre': row.antecedentes_disciplinarios_pdf_nombre,
                    'antecedentes_judiciales_pdf_nombre': row.antecedentes_judiciales_pdf_nombre,
                    'redam_pdf_nombre': row.redam_pdf_nombre,
                    'inhabilidades_sexuales_pdf_nombre': row.inhabilidades_sexuales_pdf_nombre,
                    'declaracion_capacidad_legal_pdf_nombre': row.declaracion_capacidad_legal_pdf_nombre,
                    'estado_control': row.estado_control,
                    'resultado_certificados': row.resultado_certificados,
                    'emprendimiento_formalizado': row.emprendimiento_formalizado,
                    'matricula_mercantil_pdf_nombre': row.matricula_mercantil_pdf_nombre,
                    'facturas_6meses_pdf_nombre': row.facturas_6meses_pdf_nombre,
                    'publicaciones_redes_pdf_nombre': row.publicaciones_redes_pdf_nombre,
                    'registro_ventas_pdf_nombre': row.registro_ventas_pdf_nombre,
                    'mujer_cabeza_familia': row.mujer_cabeza_familia,
                    'victima_conflicto': row.victima_conflicto,
                    'persona_discapacidad': row.persona_discapacidad,
                    'pertenencia_etnica': row.pertenencia_etnica,
                    'sisben_grupo': row.sisben_grupo,
                    'persona_reincorporacion': row.persona_reincorporacion,
                    'tiempo_funcionamiento': row.tiempo_funcionamiento,
                    'empleos_generados': row.empleos_generados,
                    'acceso_mercados': row.acceso_mercados,
                    'financiado_estado': row.financiado_estado,
                    'financiado_regalias': row.financiado_regalias,
                    'financiado_camara_comercio': row.financiado_camara_comercio,
                    'financiado_incubadoras': row.financiado_incubadoras,
                    'financiado_otro': row.financiado_otro,
                    'financiado_otro_texto': row.financiado_otro_texto,
                    'declara_veraz': row.declara_veraz,
                    'declara_no_beneficiario': row.declara_no_beneficiario,
                    'acepta_terminos': row.acepta_terminos,
                    'fecha_aceptacion_terminos': row.fecha_aceptacion_terminos.isoformat() if row.fecha_aceptacion_terminos else None,
                    'estado_inscripcion': row.estado_inscripcion,
                    'paso_actual': row.paso_actual,
                    'fecha_ultimo_guardado': row.fecha_ultimo_guardado.isoformat() if row.fecha_ultimo_guardado else None,
                    'formulario_enviado': row.formulario_enviado,
                    'rol': row.rol,
                    'fecha_creacion': row.fecha_creacion.isoformat() if row.fecha_creacion else None,
                    'fecha_actualizacion': row.fecha_actualizacion.isoformat() if row.fecha_actualizacion else None,
                    'fecha_finalizacion': row.fecha_finalizacion.isoformat() if row.fecha_finalizacion else None,
                    'estado_cuenta': row.estado_cuenta,
                    'convocatoria': row.convocatoria,
                    'emprendimiento_nombre': row.emprendimiento_nombre,
                    'emprendimiento_sector': row.emprendimiento_sector,
                    'tipo_persona': row.tipo_persona,
                    'fase_actual': 'inscripcion',  # Por defecto
                    'fecha_entrada_fase': None,
                    'fase_completada': False,
                    'documentos_subidos': doc_count
                }
                users_data.append(user_data)
            except Exception as e:
                print(f"Error procesando usuario: {e}")
                continue
        
        print(f"Pagina {page}/{total_pages} - Usuarios procesados: {len(users_data)} de {total_users} totales")
        
        # Devolver respuesta con metadatos de paginaciÃ³n
        return jsonify({
            'users': users_data,
            'pagination': {
                'page': page,
                'per_page': per_page,
                'total_users': total_users,
                'total_pages': total_pages,
                'has_next': has_next,
                'has_prev': has_prev
            }
        }), 200
        
    except Exception as e:
        print(f"Error al obtener todos los usuarios: {e}")
        import traceback
        traceback.print_exc()
        return jsonify({'error': f'Error al obtener todos los usuarios: {str(e)}'}), 500

@admin_bp.route('/users/<int:user_id>/phase', methods=['POST'])
@token_required
@admin_or_evaluador_required
def update_user_phase(current_user, user_id):
    """Actualizar la fase de un usuario especÃ­fico"""
    try:
        user = User.query.get_or_404(user_id)
        data = request.json
        nueva_fase = data.get('nueva_fase')
        
        # Validar fase vÃ¡lida
        fases_validas = ['inscripcion', 'formacion', 'entrega_activos']
        if nueva_fase not in fases_validas:
            return jsonify({'error': 'Fase no vÃ¡lida'}), 400
        
        # Obtener fase actual
        fase_actual = getattr(user, 'fase_actual', 'inscripcion')
        
        # Actualizar fase
        user.fase_actual = nueva_fase
        user.fecha_entrada_fase = datetime.utcnow()
        user.fase_completada = False
        
        db.session.commit()
        
        # Crear notificaciÃ³n de cambio de fase
        mensajes = {
            'inscripcion': 'Tu emprendimiento ha pasado a la fase de InscripciÃ³n y SelecciÃ³n',
            'formacion': 'Tu emprendimiento ha pasado a la fase de FormaciÃ³n',
            'entrega_activos': 'Tu emprendimiento ha pasado a la fase de Entrega de Activos Productivos'
        }
        
        mensaje = mensajes.get(nueva_fase, f'Tu emprendimiento ha pasado a la fase: {nueva_fase}')
        
        notificacion = Notificacion(
            user_id=user.id,
            mensaje=mensaje,
            fase_nueva=nueva_fase
        )
        db.session.add(notificacion)
        
        # Log de actividad
        db.session.add(LogActividad(
            usuario_id=user.id,
            accion='cambio_fase_admin',
            detalles=f"Administrador cambiÃ³ fase de {fase_actual} a {nueva_fase}"
        ))
        db.session.commit()
        
        return jsonify({
            'message': f'Fase actualizada exitosamente a {nueva_fase}',
            'fase_actual': nueva_fase,
            'fecha_entrada_fase': user.fecha_entrada_fase.isoformat()
        }), 200
        
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': f'Error al actualizar fase: {str(e)}'}), 500

@admin_bp.route('/export/phases', methods=['GET'])
@token_required
@admin_required
def export_phase_report(current_user):
    """Exportar reporte de fases en Excel"""
    try:
        users = User.query.all()
        
        # Crear workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Reporte de Fases"
        
        # Estilos
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")
        
        # Encabezados
        headers = [
            'ID', 'Nombre', 'Apellido', 'Email', 'TelÃ©fono', 'Municipio',
            'Emprendimiento', 'Sector', 'Tipo Persona', 'Fase Actual',
            'Fecha Entrada Fase', 'Fase Completada', 'Estado InscripciÃ³n',
            'Formulario Enviado', 'Fecha CreaciÃ³n', 'Estado Cuenta'
        ]
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
        
        # Datos
        for row, user in enumerate(users, 2):
            ws.cell(row=row, column=1, value=user.id)
            ws.cell(row=row, column=2, value=user.nombre)
            ws.cell(row=row, column=3, value=user.apellido)
            ws.cell(row=row, column=4, value=user.email)
            ws.cell(row=row, column=5, value=user.telefono)
            ws.cell(row=row, column=6, value=user.municipio)
            ws.cell(row=row, column=7, value=user.emprendimiento_nombre)
            ws.cell(row=row, column=8, value=user.emprendimiento_sector)
            ws.cell(row=row, column=9, value=user.tipo_persona)
            ws.cell(row=row, column=10, value=getattr(user, 'fase_actual', 'inscripcion'))
            ws.cell(row=row, column=11, value=user.fecha_entrada_fase.strftime('%Y-%m-%d %H:%M:%S') if hasattr(user, 'fecha_entrada_fase') and user.fecha_entrada_fase else 'N/A')
            ws.cell(row=row, column=12, value='SÃ­' if getattr(user, 'fase_completada', False) else 'No')
            ws.cell(row=row, column=13, value=user.estado_inscripcion)
            ws.cell(row=row, column=14, value='SÃ­' if user.formulario_enviado else 'No')
            ws.cell(row=row, column=15, value=user.fecha_creacion.strftime('%Y-%m-%d %H:%M:%S') if user.fecha_creacion else 'N/A')
            ws.cell(row=row, column=16, value=user.estado_cuenta)
        
        # Ajustar ancho de columnas
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 15
        
        # Guardar en archivo temporal
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx')
        wb.save(tmp.name)
        tmp.flush()
        tmp.seek(0)
        
        return send_file(
            tmp.name,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'reporte_fases_{datetime.now().strftime("%Y%m%d")}.xlsx'
        )
        
    except Exception as e:
        return jsonify({'error': f'Error al exportar reporte: {str(e)}'}), 500

# ========== ENDPOINT PARA CREAR TABLAS DE CONFIGURACIÃ“N ==========

@admin_bp.route('/create-config-tables', methods=['POST', 'GET'])
@token_required
@admin_required
def create_config_tables(current_user):
    """Crear las tablas de configuraciÃ³n dinÃ¡mica"""
    try:
        # Crear todas las tablas
        db.create_all()
        
        # Verificar que las tablas existen
        inspector = db.inspect(db.engine)
        existing_tables = inspector.get_table_names()
        
        config_tables = [
            'formulario_campos',
            'criterios_evaluacion_config', 
            'cupos_municipio_config',
            'documentos_config'
        ]
        
        created_tables = []
        for table in config_tables:
            if table in existing_tables:
                created_tables.append(table)
        
        return jsonify({
            'message': 'Tablas de configuraciÃ³n creadas exitosamente',
            'created_tables': created_tables,
            'total_tables_in_db': len(existing_tables),
            'all_tables': existing_tables
        }), 200
        
    except Exception as e:
        return jsonify({'error': f'Error creando tablas: {str(e)}'}), 500

@admin_bp.route('/test-deploy', methods=['POST'])
@token_required
@admin_required
def test_deploy_endpoint(current_user):
    """Endpoint de prueba para verificar despliegues."""
    return jsonify({'message': 'Despliegue de prueba exitoso!', 'user': current_user.email}), 200

@admin_bp.route('/config/formulario-campos', methods=['GET'])
@token_required
@admin_required
def get_formulario_campos(current_user):
    """Obtener campos del formulario"""
    try:
        campos = FormularioCampo.query.filter_by(es_activo=True).order_by(FormularioCampo.orden).all()
        return jsonify([campo.to_dict() for campo in campos])
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@admin_bp.route('/config/formulario-campos', methods=['POST'])
@token_required
@admin_required
def create_formulario_campo(current_user):
    """Crear un nuevo campo del formulario"""
    try:
        data = request.json
        campo = FormularioCampo(
            nombre_campo=data['nombre_campo'],
            etiqueta=data['etiqueta'],
            tipo_campo=data['tipo_campo'],
            es_obligatorio=data.get('es_obligatorio', False),
            es_subsanable=data.get('es_subsanable', True),
            orden=data.get('orden', 0),
            seccion=data.get('seccion', 'general'),
            opciones=data.get('opciones'),
            validaciones=data.get('validaciones'),
            descripcion=data.get('descripcion')
        )
        db.session.add(campo)
        db.session.commit()
        return jsonify({'message': 'Campo creado exitosamente', 'campo': campo.to_dict()}), 201
    except Exception as e:
        db.session.rollback()
        return jsonify({'error': str(e)}), 400

@admin_bp.route('/users/<int:user_id>/detailed-info', methods=['GET'])
@token_required
@admin_or_evaluador_required
def get_user_detailed_info(current_user, user_id):
    """Obtener informaciÃ³n completa del usuario incluyendo archivos de S3"""
    try:
        from ..services.s3_service import S3Service
        
        user = User.query.get_or_404(user_id)
        
        # InformaciÃ³n bÃ¡sica del usuario
        user_info = user.to_dict()
        
        # Obtener archivos de S3
        s3_service = S3Service()
        
        # Mapeo de tipos de documentos con nombres legibles y secciones
        document_types_mapping = {
            'obligatorios/tdr': {'name': 'TDR - TÃ©rminos y Condiciones', 'section': 'obligatorios'},
            'obligatorios/uso-imagen': {'name': 'AutorizaciÃ³n Uso de Imagen', 'section': 'obligatorios'},
            'obligatorios/plan-negocio': {'name': 'Plan de Negocio (Excel)', 'section': 'obligatorios'},
            'obligatorios/vecindad': {'name': 'Certificado de Vecindad', 'section': 'obligatorios'},
            'por-tipo/persona-natural/cedula': {'name': 'CÃ©dula de CiudadanÃ­a', 'section': 'por_tipo'},
            'por-tipo/persona-natural/rut': {'name': 'RUT Persona Natural', 'section': 'por_tipo'},
            'por-tipo/persona-juridica/camara-comercio': {'name': 'Certificado CÃ¡mara de Comercio', 'section': 'por_tipo'},
            'por-tipo/persona-juridica/rut-empresa': {'name': 'RUT Empresa', 'section': 'por_tipo'},
            'por-tipo/persona-juridica/cedula-representante': {'name': 'CÃ©dula Representante Legal', 'section': 'por_tipo'},
            'por-tipo/persona-juridica/certificado-existencia': {'name': 'Certificado de Existencia', 'section': 'por_tipo'},
            'diferenciales/ruv': {'name': 'RUV - Registro Ãšnico de VÃ­ctimas', 'section': 'diferenciales'},
            'diferenciales/sisben': {'name': 'Certificado SISBEN', 'section': 'diferenciales'},
            'diferenciales/grupo-etnico': {'name': 'Certificado Grupo Ã‰tnico', 'section': 'diferenciales'},
            'diferenciales/arn': {'name': 'Certificado ARN', 'section': 'diferenciales'},
            'diferenciales/discapacidad': {'name': 'Certificado de Discapacidad', 'section': 'diferenciales'},
            'diferenciales/mujer-cabeza-familia': {'name': 'Mujer Cabeza de Familia', 'section': 'diferenciales'},
            'diferenciales/persona-discapacidad': {'name': 'Persona en SituaciÃ³n de Discapacidad', 'section': 'diferenciales'},
            'control/antecedentes-fiscales': {'name': 'Antecedentes Fiscales', 'section': 'control'},
            'control/antecedentes-disciplinarios': {'name': 'Antecedentes Disciplinarios', 'section': 'control'},
            'control/antecedentes-judiciales': {'name': 'Antecedentes Judiciales', 'section': 'control'},
            'control/antecedentes-contraloria': {'name': 'Antecedentes ContralorÃ­a', 'section': 'control'},
            'control/antecedentes-procuraduria': {'name': 'Antecedentes ProcuradurÃ­a', 'section': 'control'},
            'control/redam': {'name': 'Certificado REDAM', 'section': 'control'},
            'control/rnmc': {'name': 'Certificado RNMC', 'section': 'control'},
            'control/inhabilidades-sexuales': {'name': 'Antecedentes Sexuales', 'section': 'control'},
            'obligatorios/declaracion-capacidad': {'name': 'DeclaraciÃ³n de Capacidad Legal', 'section': 'obligatorios'},
            'funcionamiento/matricula-mercantil': {'name': 'MatrÃ­cula Mercantil', 'section': 'funcionamiento'},
            'funcionamiento/facturas-6meses': {'name': 'Facturas 6 Meses', 'section': 'funcionamiento'},
            'funcionamiento/publicaciones-redes': {'name': 'Publicaciones en Redes', 'section': 'funcionamiento'},
            'funcionamiento/registro-ventas': {'name': 'Registro de Ventas', 'section': 'funcionamiento'},
            'funcionamiento/comprobantes-ventas': {'name': 'Comprobantes de Ventas', 'section': 'funcionamiento'},
            'funcionamiento/facturas-venta': {'name': 'Facturas de Venta', 'section': 'funcionamiento'},
            'funcionamiento/redes-sociales': {'name': 'Redes Sociales', 'section': 'funcionamiento'},
            'videos/presentacion': {'name': 'Video de PresentaciÃ³n', 'section': 'videos'}
        }
        
        # Nombres de las secciones
        section_names = {
            'obligatorios': 'ðŸ“‹ Documentos Obligatorios',
            'por_tipo': 'ðŸ‘¤ Documentos por Tipo de Persona',
            'diferenciales': 'ðŸ·ï¸ Documentos Diferenciales',
            'control': 'ðŸ” Documentos de Control',
            'funcionamiento': 'ðŸ¢ Documentos de Funcionamiento',
            'videos': 'ðŸŽ¥ Videos de PresentaciÃ³n'
        }
        
        try:
            # Listar archivos en la carpeta del usuario
            prefix = f"usuarios/{user_id}/"
            response = s3_service.s3_client.list_objects_v2(
                Bucket=s3_service.bucket_name,
                Prefix=prefix
            )
            
            # Diccionario para agrupar archivos por tipo de documento
            files_by_type = {}
            
            if 'Contents' in response:
                for obj in response['Contents']:
                    file_key = obj['Key']
                    
                    # Extraer informaciÃ³n del archivo
                    file_parts = file_key.split('/')
                    if len(file_parts) >= 4:  # usuarios/ID/documentos/tipo/archivo
                        filename = file_parts[-1]
                        
                        # Filtrar archivos .keep y otros archivos tÃ©cnicos
                        if filename == '.keep' or filename.startswith('.'):
                            continue
                            
                        doc_type_path = '/'.join(file_parts[3:-1])  # tipo/subtipo
                        
                        # Obtener informaciÃ³n del documento - solo procesar si estÃ¡ mapeado
                        doc_info = document_types_mapping.get(doc_type_path)
                        if not doc_info:
                            # Saltar archivos no mapeados (elimina la secciÃ³n "otros")
                            continue
                        
                        # Crear objeto de archivo
                        file_obj = {
                            'key': file_key,
                            'filename': filename,
                            'document_type': doc_type_path,
                            'document_name': doc_info['name'],
                            'section': doc_info['section'],
                            'size': obj['Size'],
                            'last_modified': obj['LastModified'],
                            'last_modified_iso': obj['LastModified'].isoformat()
                        }
                        
                        # Agrupar por tipo de documento y mantener solo el mÃ¡s reciente
                        if doc_type_path not in files_by_type:
                            files_by_type[doc_type_path] = file_obj
                        else:
                            # Comparar fechas y mantener el mÃ¡s reciente
                            if obj['LastModified'] > files_by_type[doc_type_path]['last_modified']:
                                files_by_type[doc_type_path] = file_obj
            
            # Generar URLs presignadas y organizar por secciones
            files_by_section = {}
            total_files = 0
            
            for doc_type, file_obj in files_by_type.items():
                try:
                    download_url = s3_service.s3_client.generate_presigned_url(
                        'get_object',
                        Params={'Bucket': s3_service.bucket_name, 'Key': file_obj['key']},
                        ExpiresIn=3600  # 1 hora
                    )
                    file_obj['download_url'] = download_url
                except Exception as e:
                    file_obj['download_url'] = None
                    print(f"Error generando URL presignada: {e}")
                
                # Remover el objeto datetime para serializaciÃ³n JSON
                file_obj_clean = {k: v for k, v in file_obj.items() if k != 'last_modified'}
                file_obj_clean['last_modified'] = file_obj['last_modified_iso']
                
                # Organizar por secciÃ³n
                section = file_obj['section']
                if section not in files_by_section:
                    files_by_section[section] = {
                        'section_name': section_names.get(section, section.title()),
                        'files': []
                    }
                
                files_by_section[section]['files'].append(file_obj_clean)
                total_files += 1
        
        except Exception as e:
            print(f"Error listando archivos de S3: {e}")
            files_by_section = {}
            total_files = 0
        
        # InformaciÃ³n adicional del formulario
        additional_info = {}
        
        # Campos especÃ­ficos del modelo User que contienen informaciÃ³n del formulario
        form_fields = [
            'emprendimiento_formalizado', 'financiado_estado', 'financiado_regalias',
            'financiado_camara_comercio', 'financiado_incubadoras', 'financiado_otro',
            'financiado_otro_texto', 'declara_veraz', 'declara_no_beneficiario',
            'acepta_terminos', 'mujer_cabeza_familia', 'victima_conflicto', 'persona_discapacidad',
            'pertenencia_etnica', 'sisben_grupo', 'persona_reincorporacion',
            'tiempo_funcionamiento', 'empleos_generados', 'acceso_mercados',
            'paso_actual', 'estado_inscripcion', 'formulario_enviado'
        ]
        
        for field in form_fields:
            if hasattr(user, field):
                value = getattr(user, field)
                # Incluir todos los valores, incluso False y 0
                additional_info[field] = value
        
        return jsonify({
            'user': user_info,
            'files_by_section': files_by_section,
            'additional_info': additional_info,
            'total_files': total_files
        }), 200
        
    except Exception as e:
        print(f"Error obteniendo informaciÃ³n detallada del usuario: {e}")
        return jsonify({'error': f'Error al obtener informaciÃ³n del usuario: {str(e)}'}), 500

# ===== ENDPOINT PARA GENERAR CONTRASEÃ‘AS MASIVAS =====
@admin_bp.route('/users/generate-passwords', methods=['POST'])
@token_required
@admin_required
def generate_passwords_csv(current_user):
    """
    Genera contraseÃ±as para una lista de IDs y devuelve un CSV
    """
    try:
        import secrets
        import string
        from werkzeug.security import generate_password_hash
        
        data = request.json
        user_ids = data.get('user_ids', [])
        
        if not user_ids:
            return jsonify({'error': 'La lista de IDs estÃ¡ vacÃ­a'}), 400
        
        def generate_temp_password(length=12):
            """Genera una contraseÃ±a temporal segura"""
            alphabet = string.ascii_letters + string.digits
            password = ''.join(secrets.choice(string.ascii_letters) for _ in range(2))
            password += ''.join(secrets.choice(string.digits) for _ in range(2))
            password += ''.join(secrets.choice(alphabet) for _ in range(length - 4))
            password_list = list(password)
            secrets.SystemRandom().shuffle(password_list)
            return ''.join(password_list)
        
        # Crear CSV en memoria
        output = io.StringIO()
        writer = csv.writer(output)
        
        # Escribir encabezados
        writer.writerow(['ID', 'Nombre', 'Apellido', 'Email', 'CÃ©dula', 'ContraseÃ±a'])
        
        success_count = 0
        error_count = 0
        errors = []
        
        # Procesar cada ID (optimizado con commits en lotes)
        batch_size = 50  # Commit cada 50 usuarios
        users_to_update = []
        
        for user_id in user_ids:
            try:
                user_id_int = int(user_id)
                user = User.query.get(user_id_int)
                
                if not user:
                    errors.append(f"ID {user_id}: Usuario no encontrado")
                    error_count += 1
                    continue
                
                # Generar nueva contraseña
                new_password = generate_temp_password()
                
                # Actualizar contraseña y estado de cuenta a "activa"
                user.set_password(new_password)
                user.estado_cuenta = 'activa'
                users_to_update.append((user, new_password))
                
                # Commit en lotes para mejorar rendimiento
                if len(users_to_update) >= batch_size:
                    try:
                        db.session.commit()
                        # Escribir en CSV después del commit exitoso
                        for user_obj, password in users_to_update:
                            writer.writerow([
                                user_obj.id,
                                user_obj.nombre,
                                user_obj.apellido,
                                user_obj.email,
                                user_obj.numero_documento,
                                password
                            ])
                            success_count += 1
                        users_to_update = []
                    except Exception as e:
                        db.session.rollback()
                        for user_obj, password in users_to_update:
                            errors.append(f"ID {user_obj.id}: {str(e)}")
                            error_count += 1
                        users_to_update = []
                
            except ValueError:
                errors.append(f"ID {user_id}: No es un número válido")
                error_count += 1
            except Exception as e:
                db.session.rollback()
                errors.append(f"ID {user_id}: {str(e)}")
                error_count += 1
        
        # Commit final para los usuarios restantes
        if users_to_update:
            try:
                db.session.commit()
                for user_obj, password in users_to_update:
                    writer.writerow([
                        user_obj.id,
                        user_obj.nombre,
                        user_obj.apellido,
                        user_obj.email,
                        user_obj.numero_documento,
                        password
                    ])
                    success_count += 1
            except Exception as e:
                db.session.rollback()
                for user_obj, password in users_to_update:
                    errors.append(f"ID {user_obj.id}: {str(e)}")
                    error_count += 1
        
        # Preparar respuesta CSV
        output.seek(0)
        csv_content = output.getvalue()
        output.close()
        
        # Crear respuesta con CSV
        response = make_response(csv_content)
        response.headers['Content-Type'] = 'text/csv; charset=utf-8'
        response.headers['Content-Disposition'] = f'attachment; filename=passwords_{datetime.now().strftime("%Y%m%d_%H%M%S")}.csv'
        
        return response
        
    except Exception as e:
        db.session.rollback()
        print(f"Error generando contraseÃ±as: {e}")
        return jsonify({'error': f'Error al generar contraseÃ±as: {str(e)}'}), 500



# ===== ENDPOINT TEMPORAL PARA ACTIVAR USUARIOS DEL LISTADO IDs.csv =====
@admin_bp.route('/users/activate-from-ids-file', methods=['POST'])
@token_required
@admin_required
def activate_users_from_ids_file(current_user):
    """
    Endpoint temporal para activar usuarios del archivo IDs.csv
    Actualiza el estado a 'activa' para los usuarios del listado
    """
    try:
        data = request.json
        user_ids = data.get('user_ids', [])
        
        if not user_ids:
            return jsonify({'error': 'No se proporcionaron IDs de usuarios'}), 400
        
        # Actualizar usuarios en lotes
        batch_size = 100
        total_updated = 0
        total_errors = 0
        errors = []
        
        for i in range(0, len(user_ids), batch_size):
            batch = user_ids[i:i + batch_size]
            
            try:
                # Actualizar estado a 'activa'
                users = User.query.filter(User.id.in_(batch)).all()
                
                for user in users:
                    user.estado_cuenta = 'activa'
                    user.fecha_actualizacion = datetime.utcnow()
                
                db.session.commit()
                total_updated += len(users)
                
            except Exception as e:
                db.session.rollback()
                errors.append(f"Lote {i//batch_size + 1}: {str(e)}")
                total_errors += len(batch)
        
        return jsonify({
            'success': True,
            'message': f'{total_updated} usuarios actualizados a estado "activa" exitosamente',
            'total_updated': total_updated,
            'total_errors': total_errors,
            'errors': errors if errors else None
        }), 200
        
    except Exception as e:
        db.session.rollback()
        print(f"Error activando usuarios: {e}")
        return jsonify({'error': f'Error al activar usuarios: {str(e)}'}), 500


# ===== ENDPOINT TEMPORAL SIN AUTENTICACION PARA ACTIVAR USUARIOS =====
@admin_bp.route('/users/activate-from-ids-file-direct', methods=['POST'])
def activate_users_from_ids_file_direct():
    """
    Endpoint temporal SIN autenticacion para activar usuarios del archivo IDs.csv
    SOLO PARA USO INTERNO - ELIMINAR DESPUES DE USAR
    """
    try:
        # Usar IDs del request (el frontend lee el archivo y envía los IDs)
        data = request.json or {}
        user_ids = data.get('user_ids', [])
        
        if not user_ids:
            return jsonify({'error': 'No se proporcionaron IDs de usuarios'}), 400
        
        if not isinstance(user_ids, list):
            return jsonify({'error': 'user_ids debe ser una lista'}), 400
        
        # Actualizar usuarios en lotes
        batch_size = 100
        total_updated = 0
        total_errors = 0
        errors = []
        
        for i in range(0, len(user_ids), batch_size):
            batch = user_ids[i:i + batch_size]
            
            try:
                # Actualizar estado a 'activa'
                users = User.query.filter(User.id.in_(batch)).all()
                
                for user in users:
                    user.estado_cuenta = 'activa'
                    user.fecha_actualizacion = datetime.utcnow()
                
                db.session.commit()
                total_updated += len(users)
                
            except Exception as e:
                db.session.rollback()
                errors.append(f"Lote {i//batch_size + 1}: {str(e)}")
                total_errors += len(batch)
        
        return jsonify({
            'success': True,
            'message': f'{total_updated} usuarios actualizados a estado "activa" exitosamente',
            'total_updated': total_updated,
            'total_processed': len(user_ids),
            'total_errors': total_errors,
            'errors': errors if errors else None
        }), 200
        
    except Exception as e:
        db.session.rollback()
        print(f"Error activando usuarios: {e}")
        return jsonify({'error': f'Error al activar usuarios: {str(e)}'}), 500
